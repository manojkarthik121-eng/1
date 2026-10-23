#!/usr/bin/env python3
"""Build the family finance workbook from a JSON plan.

  build_workbook.py plan.json --out Family_Finance.xlsx [--transactions OUT/transactions.csv]

All totals, variances, goal SIPs and net worth are live Excel formulas, so the family
can keep editing the workbook and it stays correct.

plan.json (every section optional except "family"):
{
  "family": "Sharma Family", "month": "2026-10",
  "income":   [{"source": "Rahul - salary", "monthly": 145000}],
  "budget":   [{"category": "Rent", "group": "Needs", "budget": 32000, "actual": 32000}],
  "goals":    [{"name": "Emergency fund", "target_today": 400000, "years": 1, "saved": 150000,
                "return": 0.065, "inflation": 0.0, "priority": 1}],
  "debts":    [{"name": "Credit card", "balance": 85000, "rate": 0.42, "emi": 5000}],
  "assets":   [{"name": "EPF", "type": "Retirement", "value": 640000}],
  "insurance":[{"person": "Rahul", "type": "Term", "cover": 15000000, "premium_yearly": 18000, "status": "OK"}],
  "actions":  [{"when": "This week", "action": "Stop Dream11", "impact_monthly": 4000}]
}
"""
import argparse
import csv
import json

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INR = '[>=10000000]"₹"##\\,##\\,##\\,##0;[>=100000]"₹"##\\,##\\,##0;"₹"##,##0'
PCT = "0.0%"
NAVY, TEAL, LIGHT, GREY = "1F3A5F", "2A9D8F", "EAF4F4", "F5F5F5"
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TOTAL_FILL = PatternFill("solid", fgColor=LIGHT)
INPUT_FONT = Font(color="1F4E9E")  # blue = you can edit this
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
RED_FILL = PatternFill("solid", fgColor="FBE3E1")
GREEN_FILL = PatternFill("solid", fgColor="E3F4E8")


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="666666")


def header(ws, row, names, widths=None):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def put(ws, row, col, value, fmt=None, editable=False, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if editable:
        c.font = INPUT_FONT
    if bold:
        c.font = Font(bold=True)
    return c


def total_row(ws, row, label, cols, first, last, fmt=INR, ncols=None):
    put(ws, row, 1, label, bold=True)
    for col in cols:
        L = get_column_letter(col)
        put(ws, row, col, f"=SUM({L}{first}:{L}{last})", fmt, bold=True)
    for col in range(1, (ncols or max(cols)) + 1):
        ws.cell(row=row, column=col).fill = TOTAL_FILL


def sheet_budget(wb, plan):
    ws = wb.create_sheet("Budget")
    title(ws, "Monthly Income & Budget", "Blue numbers are inputs - change them and everything recalculates.")
    r = 4
    header(ws, r, ["Income source", "Monthly"], None)
    inc = plan.get("income", [])
    first = r + 1
    for i, x in enumerate(inc):
        put(ws, first + i, 1, x["source"])
        put(ws, first + i, 2, x["monthly"], INR, editable=True)
    last = first + max(len(inc), 1) - 1
    r = last + 1
    total_row(ws, r, "Total income", [2], first, last)
    income_cell = f"Budget!$B${r}"

    r += 2
    header(ws, r, ["Category", "Group", "Budget", "Actual", "Difference", "% of income"],
           [34, 12, 15, 15, 15, 12])
    rows = plan.get("budget", [])
    bfirst = r + 1
    for i, x in enumerate(rows):
        rr = bfirst + i
        put(ws, rr, 1, x["category"])
        put(ws, rr, 2, x.get("group", ""))
        put(ws, rr, 3, x.get("budget", 0), INR, editable=True)
        put(ws, rr, 4, x.get("actual", 0), INR, editable=True)
        put(ws, rr, 5, f"=C{rr}-D{rr}", INR)
        put(ws, rr, 6, f"=IFERROR(D{rr}/{income_cell},0)", PCT)
    blast = bfirst + max(len(rows), 1) - 1
    ws.conditional_formatting.add(f"E{bfirst}:E{blast}", CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
    ws.conditional_formatting.add(f"E{bfirst}:E{blast}", CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN_FILL))
    r = blast + 1
    total_row(ws, r, "Total outflow", [3, 4, 5], bfirst, blast, ncols=6)
    put(ws, r, 6, f"=IFERROR(D{r}/{income_cell},0)", PCT, bold=True)
    tot = r

    r += 2
    put(ws, r, 1, "Split by group (actual)", bold=True)
    r += 1
    header_row = r
    for i, n in enumerate(["Group", "Actual / month", "% of income", "Healthy guide"], 1):
        c = ws.cell(row=r, column=i, value=n)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    guides = {"Needs": "≤ 50%", "Wants": "≤ 20-30%", "Debt": "EMIs ≤ 30-40%", "Savings": "≥ 20-30%"}
    gfirst = r + 1
    for i, g in enumerate(["Needs", "Wants", "Debt", "Savings"]):
        rr = gfirst + i
        put(ws, rr, 1, g)
        put(ws, rr, 2, f'=SUMIF($B${bfirst}:$B${blast},A{rr},$D${bfirst}:$D${blast})', INR)
        put(ws, rr, 3, f"=IFERROR(B{rr}/{income_cell},0)", PCT)
        put(ws, rr, 4, guides[g])
    r = gfirst + 4
    put(ws, r, 1, "Left over (unallocated)", bold=True)
    put(ws, r, 2, f"={income_cell}-D{tot}", INR, bold=True)
    ws.conditional_formatting.add(f"B{r}", CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))

    if rows:
        pie = PieChart()
        pie.title = "Where the money goes"
        pie.add_data(Reference(ws, min_col=2, min_row=gfirst, max_row=gfirst + 3), titles_from_data=False)
        pie.set_categories(Reference(ws, min_col=1, min_row=gfirst, max_row=gfirst + 3))
        pie.height, pie.width = 8, 12
        ws.add_chart(pie, "H4")
    return {"income": income_cell, "outflow": f"Budget!$D${tot}", "leftover": f"Budget!$B${r}",
            "savings": f"Budget!$B${gfirst + 3}", "debt": f"Budget!$B${gfirst + 2}",
            "wants": f"Budget!$B${gfirst + 1}", "needs": f"Budget!$B${gfirst}"}


def sheet_goals(wb, plan):
    ws = wb.create_sheet("Goals")
    title(ws, "Goals & Monthly SIP Needed",
          "Future cost = today's cost grown by inflation. SIP needed = what to invest monthly to reach it.")
    r = 4
    header(ws, r, ["Priority", "Goal", "Cost today", "Years", "Inflation", "Future cost", "Saved so far",
                   "Expected return", "SIP needed / month", "Suggested home"],
           [9, 30, 15, 8, 10, 16, 15, 11, 17, 38])
    goals = sorted(plan.get("goals", []), key=lambda g: g.get("priority", 99))
    first = r + 1
    for i, g in enumerate(goals):
        rr = first + i
        put(ws, rr, 1, g.get("priority", i + 1))
        put(ws, rr, 2, g["name"])
        put(ws, rr, 3, g["target_today"], INR, editable=True)
        put(ws, rr, 4, g["years"], "0.0", editable=True)
        put(ws, rr, 5, g.get("inflation", 0.06), PCT, editable=True)
        put(ws, rr, 6, f"=C{rr}*(1+E{rr})^D{rr}", INR)
        put(ws, rr, 7, g.get("saved", 0), INR, editable=True)
        put(ws, rr, 8, g.get("return", 0.10), PCT, editable=True)
        # SIP at start of month (type=1) so it matches fincalc.py
        put(ws, rr, 9, f"=MAX(0,-PMT((1+H{rr})^(1/12)-1,D{rr}*12,-G{rr},F{rr},1))", INR)
        put(ws, rr, 10, g.get("vehicle", ""))
        ws.cell(row=rr, column=10).alignment = Alignment(wrap_text=True)
    last = first + max(len(goals), 1) - 1
    total_row(ws, last + 1, "Total monthly SIP for all goals", [9], first, last, ncols=10)
    ws.cell(row=last + 1, column=1).value = None
    ws.cell(row=last + 1, column=2, value="Total monthly SIP for all goals").font = Font(bold=True)
    return f"Goals!$I${last + 1}"


def sheet_debts(wb, plan):
    ws = wb.create_sheet("Debts")
    title(ws, "Debts - pay off in this order",
          "Sorted highest interest first (avalanche). Put every spare rupee on row 1 until it is gone, then the next.")
    r = 4
    header(ws, r, ["Order", "Debt", "Balance", "Interest rate", "EMI / minimum", "Interest this month",
                   "Months left at EMI"], [8, 30, 15, 12, 15, 16, 16])
    debts = sorted(plan.get("debts", []), key=lambda d: -d["rate"])
    first = r + 1
    for i, d in enumerate(debts):
        rr = first + i
        put(ws, rr, 1, i + 1)
        put(ws, rr, 2, d["name"])
        put(ws, rr, 3, d["balance"], INR, editable=True)
        put(ws, rr, 4, d["rate"], PCT, editable=True)
        put(ws, rr, 5, d.get("emi", 0), INR, editable=True)
        put(ws, rr, 6, f"=C{rr}*D{rr}/12", INR)
        put(ws, rr, 7, f'=IFERROR(NPER(D{rr}/12,-E{rr},C{rr}),"never - EMI below interest")', "0")
    last = first + max(len(debts), 1) - 1
    ws.conditional_formatting.add(f"D{first}:D{last}", CellIsRule(operator="greaterThan", formula=["0.15"], fill=RED_FILL))
    total_row(ws, last + 1, "Total", [3, 5, 6], first, last, ncols=7)
    return f"Debts!$C${last + 1}", f"Debts!$F${last + 1}"


def sheet_networth(wb, plan, debt_total):
    ws = wb.create_sheet("Net Worth")
    title(ws, "Net Worth", "Update the values every 3-6 months and note the date - this is the scoreboard.")
    r = 4
    header(ws, r, ["Asset", "Type", "Value", "Share"], [32, 18, 16, 10])
    assets = plan.get("assets", [])
    first = r + 1
    for i, a in enumerate(assets):
        rr = first + i
        put(ws, rr, 1, a["name"])
        put(ws, rr, 2, a.get("type", ""))
        put(ws, rr, 3, a["value"], INR, editable=True)
    last = first + max(len(assets), 1) - 1
    tr = last + 1
    for rr in range(first, last + 1):
        put(ws, rr, 4, f"=IFERROR(C{rr}/$C${tr},0)", PCT)
    total_row(ws, tr, "Total assets", [3], first, last, ncols=4)
    put(ws, tr + 1, 1, "Total debts", bold=True)
    put(ws, tr + 1, 3, f"={debt_total}", INR, bold=True)
    put(ws, tr + 2, 1, "NET WORTH", bold=True)
    put(ws, tr + 2, 3, f"=C{tr}-C{tr + 1}", INR, bold=True)
    for col in (1, 2, 3, 4):
        ws.cell(row=tr + 2, column=col).fill = PatternFill("solid", fgColor="D8EFEA")

    r = tr + 4
    put(ws, r, 1, "Asset mix by type", bold=True)
    r += 1
    types = sorted({a.get("type", "") for a in assets})
    for i, t in enumerate(types):
        rr = r + i
        put(ws, rr, 1, t)
        put(ws, rr, 3, f'=SUMIF($B${first}:$B${last},A{rr},$C${first}:$C${last})', INR)
        put(ws, rr, 4, f"=IFERROR(C{rr}/$C${tr},0)", PCT)
    return f"'Net Worth'!$C${tr + 2}", f"'Net Worth'!$C${tr}"


def sheet_insurance(wb, plan):
    rows = plan.get("insurance", [])
    if not rows:
        return
    ws = wb.create_sheet("Insurance")
    title(ws, "Protection check", "Term cover ≈ 10-15× yearly income of each earner; health cover ≥ ₹10-25 lakh family floater.")
    header(ws, 4, ["Person", "Type", "Cover", "Premium / year", "Status / action"], [20, 20, 16, 16, 60])
    for i, x in enumerate(rows):
        rr = 5 + i
        put(ws, rr, 1, x.get("person", ""))
        put(ws, rr, 2, x.get("type", ""))
        put(ws, rr, 3, x.get("cover", 0), INR, editable=True)
        put(ws, rr, 4, x.get("premium_yearly", 0), INR, editable=True)
        put(ws, rr, 5, x.get("status", "")).alignment = Alignment(wrap_text=True)


def sheet_actions(wb, plan):
    ws = wb.create_sheet("Action Plan")
    title(ws, "Action Plan", "Tick 'Done' as you go. Review this sheet on the 1st of every month.")
    header(ws, 4, ["#", "When", "Action", "Impact / month", "Done?"], [5, 14, 80, 16, 8])
    acts = plan.get("actions", [])
    for i, a in enumerate(acts):
        rr = 5 + i
        put(ws, rr, 1, i + 1)
        put(ws, rr, 2, a.get("when", ""))
        put(ws, rr, 3, a["action"]).alignment = Alignment(wrap_text=True, vertical="top")
        put(ws, rr, 4, a.get("impact_monthly") or None, INR)
        put(ws, rr, 5, "", editable=True)
    if acts:
        last = 4 + len(acts)
        total_row(ws, last + 1, "", [4], 5, last, ncols=5)
        ws.cell(row=last + 1, column=3, value="Total monthly money freed / redirected").font = Font(bold=True)


def sheet_transactions(wb, path):
    ws = wb.create_sheet("Transactions")
    title(ws, "Categorised transactions", "Fix any wrong category here; the Category Summary sheet updates itself.")
    with open(path, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    cols = ["date", "account", "description", "merchant", "amount", "category", "group"]
    header(ws, 4, [c.title() for c in cols], [12, 14, 52, 24, 14, 28, 11])
    for i, t in enumerate(rows):
        rr = 5 + i
        for j, c in enumerate(cols, 1):
            v = float(t[c]) if c == "amount" else t[c]
            cell = ws.cell(row=rr, column=j, value=v)
            if c == "amount":
                cell.number_format = '"₹"#,##0;[Red]-"₹"#,##0'
            if c in ("category", "group"):
                cell.font = INPUT_FONT
    last = 4 + len(rows)
    ws.auto_filter.ref = f"A4:G{last}"

    cs = wb.create_sheet("Category Summary")
    title(cs, "Spending by category (from Transactions)")
    header(cs, 4, ["Category", "Group", "Total in period"], [30, 12, 18])
    cats = {}
    for t in rows:
        cats.setdefault(t["category"], t["group"])
    for i, (c, g) in enumerate(sorted(cats.items(), key=lambda kv: kv[1])):
        rr = 5 + i
        put(cs, rr, 1, c)
        put(cs, rr, 2, g)
        put(cs, rr, 3, f"=-SUMIF(Transactions!$F$5:$F${last},A{rr},Transactions!$E$5:$E${last})", INR)


def sheet_summary(wb, plan, refs):
    ws = wb["Summary"]
    title(ws, f"{plan['family']} - Money Dashboard", f"Month: {plan.get('month', '')}")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 50
    items = [
        ("Monthly income", f"={refs['income']}", INR, ""),
        ("Monthly outflow (budget actuals)", f"={refs['outflow']}", INR, ""),
        ("Unallocated each month", f"={refs['leftover']}", INR, "Should be ~0: give every rupee a job"),
        ("Savings & investing rate", f"=IFERROR({refs['savings']}/{refs['income']},0)", PCT, "Aim for 20% now, 30%+ over time"),
        ("Wants as % of income", f"=IFERROR({refs['wants']}/{refs['income']},0)", PCT, "Guide: ≤ 20-30%"),
        ("EMIs as % of income", f"=IFERROR({refs['debt']}/{refs['income']},0)", PCT, "Danger above 40%"),
        ("Total debt", f"={refs['debt_total']}", INR, ""),
        ("Interest paid this month", f"={refs['interest']}", INR, "Money that buys nothing"),
        ("SIP needed for all goals", f"={refs['goal_sip']}", INR, "Compare with what you invest today"),
        ("Total assets", f"={refs['assets']}", INR, ""),
        ("NET WORTH", f"={refs['networth']}", INR, "The number to grow every year"),
    ]
    header(ws, 4, ["Metric", "Value", "Guide"])
    for i, (k, f, fmt, note) in enumerate(items):
        rr = 5 + i
        put(ws, rr, 1, k, bold=(k == "NET WORTH"))
        put(ws, rr, 2, f, fmt, bold=(k == "NET WORTH"))
        put(ws, rr, 3, note)
    rr = 5 + len(items) + 1
    ws.cell(row=rr, column=1, value="How to use: blue cells are inputs. Update Budget actuals monthly and "
                                    "Net Worth quarterly. Formulas do the rest.").font = Font(italic=True, color="666666")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("plan")
    ap.add_argument("--out", required=True)
    ap.add_argument("--transactions")
    a = ap.parse_args()
    plan = json.load(open(a.plan, encoding="utf-8"))

    wb = Workbook()
    wb.active.title = "Summary"
    refs = sheet_budget(wb, plan)
    refs["goal_sip"] = sheet_goals(wb, plan)
    refs["debt_total"], refs["interest"] = sheet_debts(wb, plan)
    refs["networth"], refs["assets"] = sheet_networth(wb, plan, refs["debt_total"])
    sheet_insurance(wb, plan)
    sheet_actions(wb, plan)
    if a.transactions:
        sheet_transactions(wb, a.transactions)
    sheet_summary(wb, plan, refs)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.save(a.out)
    print(f"saved {a.out} with sheets: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
