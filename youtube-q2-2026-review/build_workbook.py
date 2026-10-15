"""Builds the Q2 2026 YouTube Performance Review workbook.

    python build_workbook.py                 -> YouTube_Q2_2026_Review.xlsx (blank, ready for your data)
    python build_workbook.py --demo          -> YouTube_Q2_2026_Review_SAMPLE.xlsx (fictional example data)

Every number on the Dashboard is a live formula driven by the yellow input cells,
so the workbook recalculates when you paste in your YouTube Studio export.
"""
import random
import sys
from datetime import date, time, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- styling
FONT = "Arial"
NAVY, ACCENT, LIGHT, PALE = "1F3864", "2E75B6", "DDEBF7", "F3F8FC"
INPUT_FILL, INPUT_FONT = "FFF4C2", "0000FF"
CALC_HDR, GREY_TXT = "595959", "7F7F7F"
RED_BG, RED_TX = "FFC7CE", "9C0006"
AMB_BG, AMB_TX = "FFEB9C", "9C5700"
GRN_BG, GRN_TX = "C6EFCE", "006100"
ELT_BG, ELT_TX = "00B050", "FFFFFF"
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
PCT = '0.0"%"'          # all percentages are stored as percentage points (5.2 = 5.2%)
INT = "#,##0"
MONEY = '"$"#,##0.00'


def f(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def put(ws, ref, value, size=10, bold=False, color="000000", bg=None, align=None, fmt=None, border=False, italic=False):
    c = ws[ref]
    c.value = value
    c.font = f(size, bold, color, italic)
    if bg:
        c.fill = fill(bg)
    if align:
        c.alignment = align
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BOX
    return c


def merge_put(ws, rng, value, **kw):
    ws.merge_cells(rng)
    first = rng.split(":")[0]
    c = put(ws, first, value, **kw)
    if kw.get("border"):
        for row in ws[rng]:
            for cell in row:
                cell.border = BOX
    if kw.get("bg"):
        for row in ws[rng]:
            for cell in row:
                cell.fill = fill(kw["bg"])
    return c


def banner(ws, rng, title, subtitle_ref=None, subtitle=None):
    merge_put(ws, rng, title, size=18, bold=True, color="FFFFFF", bg=NAVY,
              align=Alignment(horizontal="left", vertical="center", indent=1))
    ws.row_dimensions[int("".join(ch for ch in rng.split(":")[0] if ch.isdigit()))].height = 36
    if subtitle_ref:
        put(ws, subtitle_ref, subtitle, size=10, italic=True, color=GREY_TXT)


def section(ws, rng, text):
    merge_put(ws, rng, text, size=12, bold=True, color="FFFFFF", bg=ACCENT,
              align=Alignment(horizontal="left", vertical="center", indent=1))
    ws.row_dimensions[int("".join(ch for ch in rng.split(":")[0] if ch.isdigit()))].height = 22


def header_row(ws, row, labels, start_col=1, bg=NAVY, height=32):
    for i, label in enumerate(labels):
        col = get_column_letter(start_col + i)
        put(ws, f"{col}{row}", label, bold=True, color="FFFFFF", bg=bg, align=CENTER, border=True)
    ws.row_dimensions[row].height = height


def grade_rules(ws, rng, anchor):
    """Traffic-light colouring for any cell whose text contains a grade word."""
    for word, bg, tx in [("Needs work", RED_BG, RED_TX), ("Average", AMB_BG, AMB_TX),
                         ("Strong", GRN_BG, GRN_TX), ("Elite", ELT_BG, ELT_TX)]:
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'ISNUMBER(SEARCH("{word}",{anchor}))'], fill=fill(bg), font=Font(name=FONT, bold=True, color=tx)))


def change_rules(ws, rng, anchor):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({anchor}),{anchor}>0)'],
                                                   fill=fill(GRN_BG), font=Font(name=FONT, color=GRN_TX, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({anchor}),{anchor}<0)'],
                                                   fill=fill(RED_BG), font=Font(name=FONT, color=RED_TX, bold=True)))


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------- constants
FIRST, LAST = 5, 204            # Video Data rows
VD = "'Video Data'"


def R(col):
    return f"{VD}!${col}${FIRST}:${col}${LAST}"


TRAFFIC_SOURCES = [
    ("Browse features", "Home feed & subscriptions. Grows when viewers are satisfied — the algorithm trusts you."),
    ("Suggested videos", "Shown next to/after other videos. Grows when you make content adjacent to bigger videos."),
    ("YouTube search", "People typing a query. Evergreen 'how to' AI tutorials win here and keep earning for months."),
    ("Shorts feed", "Swipe feed. Big reach, low intent — use it as a funnel to long-form."),
    ("Channel pages", "Viewers browsing your channel — a sign of loyal fans. Good playlists/trailer help."),
    ("External", "Links from Google, X, Reddit, newsletters, AI chatbots. Shows off-platform demand."),
    ("Notifications", "Bell subscribers. Reflects how engaged your core audience is."),
    ("Playlists", "Binge sessions. Organised series increase this and session time."),
    ("End screens", "Your own end screens. Low = end screens aren't compelling or are missing."),
    ("Direct or unknown", "Typed URLs, some apps/embeds. Usually small."),
    ("Other YouTube features", "Misc. surfaces (e.g. Hype, Explore)."),
    ("Hashtag pages", "Viewers browsing a hashtag. Usually small."),
]
VIEWER_TYPES = ["New viewers", "Casual viewers", "Regular viewers"]
DEVICES = ["Mobile phone", "Computer", "TV", "Tablet", "Game console"]
AGES = ["13–17", "18–24", "25–34", "35–44", "45–54", "55–64", "65+"]

# Benchmarks: (label, unit, average-from, strong-from, elite-from, meaning, source, confidence)
BENCH = [
    ("Impressions click-through rate (long-form)", "% of impressions", 4, 6, 9,
     "Packaging power: how often a shown thumbnail + title gets clicked. Naturally falls as YouTube shows a video to colder audiences, so judge it together with impressions.",
     "YouTube Help cites 2–10% for most channels; 2026 cross-niche analyses put the organic average at ~4–6% with top performers 9%+. subsub.io/blog/youtube-ctr-what-is-a-good-click-through-rate ; focus-digital.co/average-youtube-ctr-organic-paid-benchmarks-2026",
     "Medium"),
    ("Average % viewed (long-form)", "% of video", 25, 35, 50,
     "Retention: how much of each video people watch. The strongest signal of whether the video delivered on its promise.",
     "vidIQ (Jun 2026): ~30% is good for long-form; platform average low-to-mid 20s; few videos exceed 50%. vidiq.com/blog/post/increase-audience-retention-youtube",
     "Medium"),
    ("Average % viewed (Shorts)", "% of Short", 50, 70, 90,
     "Shorts retention (loops count, so >100% is possible). Below ~50% means viewers swipe away early — a hook problem.",
     "vidIQ reports ~74% avg % viewed on its own Shorts; other analyses cite ~50% for 50–60s Shorts. vidiq.com/blog/post/increase-audience-retention-youtube",
     "Medium"),
    ("Engagement rate ((likes + comments) ÷ views)", "% of views", 2.5, 3.5, 5,
     "How strongly viewers react. Feeds satisfaction signals and community.",
     "OverseerOS 2026 study of 2,251 videos: median 3.48% long-form, 75th percentile ~4.9%, ~5% = top quartile. overseeros.com/blog/youtube-engagement-rate-benchmarks",
     "Medium"),
    ("Comments per 1,000 views", "comments", 1, 3, 6,
     "Conversation depth. Comments are a strong community signal and a free source of video ideas.",
     "Rule of thumb from creator practice (no reliable published benchmark). Adjust to your niche.",
     "Low — strategist estimate"),
    ("Net subscribers per 1,000 views", "subs", 2, 5, 10,
     "How well views convert into long-term audience.",
     "Rule of thumb from creator practice (no reliable published benchmark). Adjust to your niche.",
     "Low — strategist estimate"),
    ("Returning viewers share (casual + regular)", "% of viewers", 20, 35, 50,
     "Loyalty. A rising share means you're building an audience, not just getting hits.",
     "Rule of thumb from creator practice. YouTube's 2026 shift to satisfaction signals rewards repeat viewers.",
     "Low — strategist estimate"),
]
BROW = {k: 5 + i for i, k in enumerate(["ctr", "apv_long", "apv_short", "eng", "cpk", "spk", "ret"])}
MIN_IMPR_REF = "Benchmarks!$C$15"


def grade(x, r):
    return (f'IF({x}>=Benchmarks!$E${r},"Elite",IF({x}>=Benchmarks!$D${r},"Strong",'
            f'IF({x}>=Benchmarks!$C${r},"Average","Needs work")))')


# ======================================================================= sheets
def build_start(wb, demo):
    ws = wb.active
    ws.title = "Start Here"
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 2, "B": 26, "C": 16, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14, "I": 14, "J": 14})
    banner(ws, "B1:J1", "Q2 2026 YouTube Performance Review  ·  Apr 1 – Jun 30, 2026")
    put(ws, "B2", "Built as a living workbook: paste your YouTube Studio numbers into the yellow cells and every chart, grade, colour and takeaway updates automatically.",
        italic=True, color=GREY_TXT)
    if demo:
        merge_put(ws, "B3:J3", "SAMPLE FILE — every number in this workbook is FICTIONAL example data for a made-up AI channel. Use the blank workbook for your real numbers.",
                  bold=True, color=RED_TX, bg=RED_BG, align=CENTER)
        ws.row_dimensions[3].height = 30

    put(ws, "B5", "Data status", bold=True)
    merge_put(ws, "C5:J5",
              f'=IF(COUNTA({R("A")})=0,"No video data yet — follow Step 1 below.","✔ "&COUNTA({R("A")})&" videos loaded. Open the Dashboard tab.")',
              bold=True, color=NAVY, bg=LIGHT, align=Alignment(vertical="center", indent=1))

    row = 7
    section(ws, f"B{row}:J{row}", "How to use this workbook (about 15 minutes)")
    steps = [
        ("STEP 1 — Per-video data (most important)",
         "YouTube Studio → Analytics → ADVANCED MODE → 'Content' tab → date range: Custom, Apr 1 2026 – Jun 30 2026. "
         "Click the '+' / metrics picker and add: Impressions, Impressions click-through rate, Average view duration, Average percentage viewed, "
         "Watch time (hours), Likes, Comments added, Shares, Subscribers gained, Subscribers lost, Estimated revenue (if monetised). "
         "Click the download icon (top right) → 'Comma-separated values (.csv)'."),
        ("STEP 1b — Load it",
         "Easiest: run  python import_studio_export.py YouTube_Q2_2026_Review.xlsx \"Table data.csv\"  (maps columns automatically). "
         "Or copy each column from the CSV into the matching yellow column on the 'Video Data' tab (headers use YouTube's own names). "
         "Format column: type Long or Short, or leave blank (videos ≤ 3 min are treated as Shorts)."),
        ("STEP 2 — Monthly totals",
         "Studio → Analytics → Overview. Set the date range to April 2026 and copy the channel totals into the 'Monthly Trend' tab; repeat for May and June. "
         "Optional but recommended: also enter Q1 2026 (Jan 1 – Mar 31) to see quarter-over-quarter growth."),
        ("STEP 3 — Traffic & audience",
         "Advanced mode, Apr 1 – Jun 30 2026: 'Traffic source' tab → views per source; Audience → New / Casual / Regular viewers; "
         "'Device type' and 'Viewer age' tabs → views. Enter them in the 'Traffic & Audience' tab."),
        ("STEP 4 — Read & act",
         "Dashboard → your Q2 story at a glance. Video Data → every video graded and diagnosed. "
         "Insights & Research → what's changing on YouTube and in AI. Q3 Action Plan → what to do. Q3 Scorecard → targets to track."),
        ("Timing note",
         "Today Q3 (Jul–Sep 2026) has already closed. Use the Q3 Scorecard now to grade Q3 against Q2, then carry the open action items into Q4."),
    ]
    row += 1
    for title, body in steps:
        put(ws, f"B{row}", title, bold=True, color=NAVY, align=WRAP)
        merge_put(ws, f"C{row}:J{row}", body, align=WRAP)
        ws.row_dimensions[row].height = 58
        row += 1

    row += 1
    section(ws, f"B{row}:J{row}", "Colour legend")
    row += 1
    legend = [
        (INPUT_FILL, INPUT_FONT, "Yellow cell, blue text", "You type or paste here. These are the ONLY cells you need to edit."),
        ("D9D9D9", "000000", "Grey header", "Calculated automatically — please don't type over these."),
        (RED_BG, RED_TX, "Needs work", "Below typical creators — a priority for Q3."),
        (AMB_BG, AMB_TX, "Average", "In the normal range — room to grow."),
        (GRN_BG, GRN_TX, "Strong", "Above typical — protect and repeat what's working."),
        (ELT_BG, ELT_TX, "Elite", "Top-tier — this is a signature strength."),
    ]
    for bg, tx, name, desc in legend:
        put(ws, f"B{row}", name, bold=True, color=tx, bg=bg, align=CENTER, border=True)
        merge_put(ws, f"C{row}:J{row}", desc, align=Alignment(vertical="center"))
        row += 1

    row += 1
    section(ws, f"B{row}:J{row}", "Example of one row on the Video Data tab (format guide only — not real data)")
    row += 1
    ex_hdr = ["Video title", "Publish date", "Format", "Duration (sec)", "Views", "Impressions", "CTR (%)", "Avg view duration", "Avg % viewed"]
    ex_val = ["I tested 5 AI agents for a week", date(2026, 4, 14), "Long", 754, 18432, 241000, 5.8, time(0, 3, 51), 30.6]
    ex_fmt = [None, "dd-mmm-yyyy", None, INT, INT, INT, PCT, "[m]:ss", PCT]
    for i, (h, v, fm) in enumerate(zip(ex_hdr, ex_val, ex_fmt)):
        col = get_column_letter(2 + i)
        put(ws, f"{col}{row}", h, bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
        put(ws, f"{col}{row + 1}", v, color=GREY_TXT, italic=True, fmt=fm, border=True, align=CENTER)
    ws.row_dimensions[row].height = 30
    put(ws, f"B{row + 2}", "Percentages are typed exactly as YouTube shows them: 5.8 means 5.8% (not 0.058). Average view duration can be pasted as 3:51 or 0:03:51.",
        italic=True, color=GREY_TXT)

    row += 4
    section(ws, f"B{row}:J{row}", "Tabs in this workbook")
    tabs = [
        ("Dashboard", "One-page Q2 summary: KPI tiles, auto-written takeaways, benchmark health check, long vs Shorts, top & bottom videos, timing patterns, charts."),
        ("Video Data", "One row per video. Paste here. Each video gets graded on CTR, retention and engagement, scored 1–4 and diagnosed."),
        ("Monthly Trend", "April / May / June channel totals, month-over-month change and Q2 vs Q1."),
        ("Traffic & Audience", "Where views came from and who watched (new vs returning, device, age)."),
        ("Benchmarks", "The thresholds behind every grade, with sources. Edit them if your niche differs."),
        ("Insights & Research", "Research on YouTube's 2026 changes, the creator economy and the AI space — and what each means for you."),
        ("Q3 Action Plan", "Prioritised action items with owners, timing, KPI and status tracking."),
        ("Q3 Scorecard", "Q2 baseline → Q3 targets → Q3 actuals, colour-coded on/off track."),
    ]
    row += 1
    for name, desc in tabs:
        c = put(ws, f"B{row}", name, bold=True, color=ACCENT, align=Alignment(vertical="center"))
        c.hyperlink = f"#'{name}'!A1"
        merge_put(ws, f"C{row}:J{row}", desc, align=Alignment(wrap_text=True, vertical="center"))
        ws.row_dimensions[row].height = 28
        row += 1
    return ws


def build_video_data(wb):
    ws = wb.create_sheet("Video Data")
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:P1", "Video Data — one row per video (Q2 2026: Apr 1 – Jun 30)")
    put(ws, "A2", "Paste/type into the yellow columns A–P (YouTube Studio → Analytics → Advanced mode → Content, date range Apr 1 – Jun 30 2026). "
                  "Grey columns Q onward calculate automatically. Percentages as shown in Studio (5.8 = 5.8%).",
        italic=True, color=GREY_TXT)
    merge_put(ws, "A3:P3", "YOUR INPUTS  (from YouTube Studio)", bold=True, color=NAVY, bg=INPUT_FILL, align=CENTER)
    merge_put(ws, "Q3:AF3", "AUTO-CALCULATED  (do not edit)", bold=True, color="FFFFFF", bg=CALC_HDR, align=CENTER)

    inputs = [
        ("Video title", 44, None), ("Video publish time", 13, "dd-mmm-yyyy"), ("Format\n(Long / Short)", 10, None),
        ("Duration (sec)", 10, INT), ("Views", 11, INT), ("Impressions", 12, INT),
        ("Impressions click-through rate (%)", 12, PCT), ("Average view duration", 11, "[m]:ss"),
        ("Average percentage viewed (%)", 12, PCT), ("Watch time (hours)", 11, "#,##0.0"),
        ("Likes", 9, INT), ("Comments added", 10, INT), ("Shares", 9, INT),
        ("Subscribers gained", 11, INT), ("Subscribers lost", 11, INT), ("Estimated revenue (USD)", 11, MONEY),
    ]
    calcs = [
        ("Format used", 9, None), ("Length (min)", 8, "0.0"), ("Published in Q2?", 9, None), ("Publish day", 8, None),
        ("Long-form length bucket", 12, None), ("Engagement rate (%)", 11, '0.00"%"'), ("Comments per 1k views", 10, "0.0"),
        ("Net subscribers", 10, INT), ("Net subs per 1k views", 10, "0.0"),
        ("CTR grade", 12, None), ("Retention grade", 12, None), ("Engagement grade", 12, None),
        ("Video score (1–4)", 9, "0.0"), ("Diagnosis — what to do", 46, None),
        ("Rank within format", 9, None), ("CTR (eligible long-form)", 9, PCT),
    ]
    helpers = [("CTR pts", 6, "0"), ("Ret pts", 6, "0"), ("Eng pts", 6, "0")]
    hr = 4
    for i, (label, w, fm) in enumerate(inputs + calcs + helpers):
        col = get_column_letter(1 + i)
        ws.column_dimensions[col].width = w
        bg = NAVY if i < len(inputs) else (CALC_HDR if i < len(inputs) + len(calcs) else "A6A6A6")
        put(ws, f"{col}{hr}", label, bold=True, color="FFFFFF", bg=bg, align=CENTER, border=True)
    ws.row_dimensions[hr].height = 46

    for r in range(FIRST, LAST + 1):
        for i, (_, _, fm) in enumerate(inputs):
            c = ws.cell(row=r, column=1 + i)
            c.fill = fill(INPUT_FILL)
            c.font = f(10, color=INPUT_FONT)
            c.border = BOX
            if fm:
                c.number_format = fm
        fx = {
            "Q": f'=IF(A{r}="","",IF(C{r}<>"",C{r},IF(AND(D{r}<>"",D{r}<=180),"Short","Long")))',
            "R": f'=IF(OR(A{r}="",D{r}=""),"",D{r}/60)',
            "S": f'=IF(OR(A{r}="",B{r}=""),"",IF(AND(B{r}>=DATE(2026,4,1),B{r}<=DATE(2026,6,30)),"Yes","No"))',
            "T": f'=IF(OR(A{r}="",B{r}=""),"",CHOOSE(WEEKDAY(B{r}),"Sun","Mon","Tue","Wed","Thu","Fri","Sat"))',
            "U": f'=IF(OR(Q{r}<>"Long",D{r}=""),"",IF(D{r}<480,"Under 8 min",IF(D{r}<900,"8–15 min",IF(D{r}<1500,"15–25 min","25+ min"))))',
            "V": f'=IF(OR(A{r}="",N(E{r})=0),"",(N(K{r})+N(L{r}))/E{r}*100)',
            "W": f'=IF(OR(A{r}="",N(E{r})=0),"",N(L{r})/E{r}*1000)',
            "X": f'=IF(A{r}="","",N(N{r})-N(O{r}))',
            "Y": f'=IF(OR(A{r}="",N(E{r})=0),"",X{r}/E{r}*1000)',
            "Z": f'=IF(OR(A{r}="",G{r}=""),"",IF(Q{r}="Short","n/a (Short)",{grade(f"G{r}", BROW["ctr"])}))',
            "AA": f'=IF(OR(A{r}="",I{r}=""),"",IF(Q{r}="Short",{grade(f"I{r}", BROW["apv_short"])},{grade(f"I{r}", BROW["apv_long"])}))',
            "AB": f'=IF(V{r}="","",{grade(f"V{r}", BROW["eng"])})',
            "AC": f'=IF(COUNT(AG{r}:AI{r})=0,"",AVERAGE(AG{r}:AI{r}))',
            "AD": (f'=IF(A{r}="","",IF(Q{r}="Short",'
                   f'IF(AH{r}="","Add % viewed to diagnose",'
                   f'IF(AND(AH{r}>=3,N(AI{r})>=3),"Winner Short — expand it into a long-form video",'
                   f'IF(AH{r}<=1,"Hook problem — viewers swipe away in the first seconds",'
                   f'IF(AI{r}=1,"Watched but no reaction — end on a question or opinion","Solid Short — keep iterating")))),'
                   f'IF(OR(AG{r}="",AH{r}=""),"Add CTR & % viewed to diagnose",'
                   f'IF(AND(AG{r}>=3,AH{r}>=3),"WINNER — make a sequel, series or Shorts from it",'
                   f'IF(AND(AG{r}>=3,AH{r}<=2),"Click but leave — intro too slow or title over-promises",'
                   f'IF(AND(AG{r}<=2,AH{r}>=3),"Great video, weak packaging — A/B test title + thumbnail",'
                   f'IF(AND(AG{r}=1,AH{r}=1),"Topic miss — don\'t repeat this angle","Average — sharpen the hook and the title")))))))'),
            "AE": (f'=IF(OR(A{r}="",N(E{r})=0),"",Q{r}&"-"&(COUNTIFS($Q${FIRST}:$Q${LAST},Q{r},$E${FIRST}:$E${LAST},">"&E{r})'
                   f'+COUNTIFS($Q${FIRST}:Q{r},Q{r},$E${FIRST}:E{r},E{r})))'),
            "AF": f'=IF(AND(Q{r}="Long",N(F{r})>={MIN_IMPR_REF},G{r}<>""),G{r},"")',
            "AG": f'=IF(Z{r}="Elite",4,IF(Z{r}="Strong",3,IF(Z{r}="Average",2,IF(Z{r}="Needs work",1,""))))',
            "AH": f'=IF(AA{r}="Elite",4,IF(AA{r}="Strong",3,IF(AA{r}="Average",2,IF(AA{r}="Needs work",1,""))))',
            "AI": f'=IF(AB{r}="Elite",4,IF(AB{r}="Strong",3,IF(AB{r}="Average",2,IF(AB{r}="Needs work",1,""))))',
        }
        for col, formula in fx.items():
            c = ws[f"{col}{r}"]
            c.value = formula
            c.font = f(10, color=GREY_TXT if col in ("AE", "AF", "AG", "AH", "AI") else "000000")
            c.border = BOX
            c.alignment = Alignment(vertical="center", wrap_text=(col == "AD"))
        for i, (_, _, fm) in enumerate(calcs + helpers):
            if fm:
                ws.cell(row=r, column=len(inputs) + 1 + i).number_format = fm
        ws[f"A{r}"].alignment = Alignment(vertical="center", wrap_text=True)

    dv = DataValidation(type="list", formula1='"Long,Short"', allow_blank=True)
    dv.error, dv.errorTitle = "Choose Long or Short (or leave blank to auto-detect).", "Format"
    ws.add_data_validation(dv)
    dv.add(f"C{FIRST}:C{LAST}")

    rng = f"Z{FIRST}:AB{LAST}"
    grade_rules(ws, rng, f"Z{FIRST}")
    for word, bg, tx in [("WINNER", ELT_BG, ELT_TX), ("Winner Short", ELT_BG, ELT_TX),
                         ("weak packaging", AMB_BG, AMB_TX), ("Click but leave", AMB_BG, AMB_TX),
                         ("Topic miss", RED_BG, RED_TX), ("Hook problem", RED_BG, RED_TX),
                         ("no reaction", AMB_BG, AMB_TX)]:
        ws.conditional_formatting.add(f"AD{FIRST}:AD{LAST}", FormulaRule(
            formula=[f'ISNUMBER(SEARCH("{word}",AD{FIRST}))'], fill=fill(bg), font=Font(name=FONT, color=tx, bold=True)))
    ws.conditional_formatting.add(f"AC{FIRST}:AC{LAST}", ColorScaleRule(
        start_type="num", start_value=1, start_color="F8696B", mid_type="num", mid_value=2.5,
        mid_color="FFEB84", end_type="num", end_value=4, end_color="63BE7B"))
    ws.conditional_formatting.add(f"E{FIRST}:E{LAST}", DataBarRule(start_type="min", end_type="max", color="5B9BD5"))
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:AI{LAST}"
    ws["G4"].comment = Comment("Type the number exactly as Studio shows it: 5.8 means 5.8%.", "Review")
    ws["AF4"].comment = Comment("Long-form videos with at least the minimum impressions set on the Benchmarks tab (default 1,000) — used to find your best-packaged video fairly.", "Review")
    ws.column_dimensions.group("AG", "AI", hidden=True)
    return ws


def build_monthly(wb):
    ws = wb.create_sheet("Monthly Trend")
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:M1", "Monthly Trend — channel totals (YouTube Studio → Analytics → Overview)")
    put(ws, "A2", "Enter April, May, June (and optionally Q1 2026) channel totals in the yellow cells. Q2 totals, growth and video counts are calculated.",
        italic=True, color=GREY_TXT)
    cols = ["Period", "Views", "Watch time (hours)", "Subscribers gained", "Subscribers lost", "Net subscribers",
            "Impressions", "Impressions CTR (%)", "Average view duration", "Unique viewers", "Estimated revenue (USD)",
            "Long-form videos published", "Shorts published"]
    widths(ws, {"A": 24, **{get_column_letter(i): 14 for i in range(2, 14)}})
    header_row(ws, 4, cols, height=40)
    fmts = [None, INT, "#,##0.0", INT, INT, INT, INT, PCT, "[m]:ss", INT, MONEY, "0", "0"]
    months = [("April 2026", date(2026, 4, 1), date(2026, 5, 1)), ("May 2026", date(2026, 5, 1), date(2026, 6, 1)),
              ("June 2026", date(2026, 6, 1), date(2026, 7, 1))]
    for i, (label, start, end) in enumerate(months):
        r = 5 + i
        put(ws, f"A{r}", label, bold=True, border=True)
        for ci in range(2, 14):
            col = get_column_letter(ci)
            c = ws[f"{col}{r}"]
            c.border, c.number_format = BOX, fmts[ci - 1]
            if col == "F":
                c.value, c.font = f"=IF(AND(D{r}=\"\",E{r}=\"\"),\"\",N(D{r})-N(E{r}))", f(10)
            elif col in ("L", "M"):
                crit = '"<>Short"' if col == "L" else '"Short"'
                c.value = (f'=COUNTIFS({R("B")},">="&DATE({start.year},{start.month},1),{R("B")},"<"&DATE({end.year},{end.month},1),'
                           f'{R("Q")},{crit})')
                c.font = f(10)
            else:
                c.fill, c.font = fill(INPUT_FILL), f(10, color=INPUT_FONT)
    # Q2 total
    put(ws, "A8", "Q2 2026 TOTAL", bold=True, color="FFFFFF", bg=NAVY, border=True)
    for ci in range(2, 14):
        col = get_column_letter(ci)
        if col == "H":
            v = '=IF(SUM(G5:G7)=0,"",SUMPRODUCT(G5:G7,H5:H7)/SUM(G5:G7))'
        elif col == "I":
            v = '=IF(SUMPRODUCT(B5:B7,I5:I7)=0,"",SUMPRODUCT(B5:B7,I5:I7)/SUM(B5:B7))'
        else:
            v = f"=SUM({col}5:{col}7)"
        put(ws, f"{col}8", v, bold=True, bg=LIGHT, border=True, fmt=fmts[ci - 1])
    ws["J8"].comment = Comment("Approximation: monthly unique viewers overlap. For the exact Q2 figure, set Studio's date range to Apr 1 – Jun 30.", "Review")
    ws["H8"].comment = Comment("Weighted by impressions, so a big month counts more.", "Review")
    # Q1 comparison
    put(ws, "A9", "Q1 2026 (Jan–Mar) — optional", bold=True, border=True)
    for ci in range(2, 14):
        col = get_column_letter(ci)
        c = ws[f"{col}9"]
        c.border, c.number_format = BOX, fmts[ci - 1]
        if col == "F":
            c.value, c.font = '=IF(AND(D9="",E9=""),"",N(D9)-N(E9))', f(10)
        elif col in ("L", "M"):
            c.value, c.font = "", f(10)
            c.fill, c.font = fill(INPUT_FILL), f(10, color=INPUT_FONT)
        else:
            c.fill, c.font = fill(INPUT_FILL), f(10, color=INPUT_FONT)
    put(ws, "A10", "Q2 vs Q1 change (%)", bold=True, border=True)
    for ci in range(2, 14):
        col = get_column_letter(ci)
        put(ws, f"{col}10", f'=IF(OR(N({col}9)=0,{col}8=""),"",({col}8-{col}9)/ABS({col}9)*100)', border=True, fmt='+0.0"%";-0.0"%";0"%"')
    change_rules(ws, "B10:M10", "B10")

    put(ws, "A12", "May vs April (%)", bold=True, border=True)
    put(ws, "A13", "June vs May (%)", bold=True, border=True)
    put(ws, "A14", "June vs April (%)", bold=True, border=True)
    for ci in range(2, 14):
        col = get_column_letter(ci)
        for r, (a, b) in zip((12, 13, 14), ((6, 5), (7, 6), (7, 5))):
            put(ws, f"{col}{r}", f'=IF(OR(N({col}{b})=0,{col}{a}=""),"",({col}{a}-{col}{b})/ABS({col}{b})*100)',
                border=True, fmt='+0.0"%";-0.0"%";0"%"')
    change_rules(ws, "B12:M14", "B12")
    put(ws, "A16", "Green = growth, red = decline. Note: from Aug 24 2026 YouTube counts a view from the first frame, so compare Q3+ views with care (see Insights).",
        italic=True, color=GREY_TXT)
    ws.freeze_panes = "B5"
    return ws


def build_traffic(wb):
    ws = wb.create_sheet("Traffic & Audience")
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:G1", "Traffic & Audience — Q2 2026")
    put(ws, "A2", "Studio → Analytics → Advanced mode, Apr 1 – Jun 30 2026. Enter views per row in the yellow cells; shares and changes calculate.",
        italic=True, color=GREY_TXT)
    widths(ws, {"A": 26, "B": 13, "C": 13, "D": 13, "E": 12, "F": 12, "G": 70})
    header_row(ws, 4, ["Traffic source", "Q2 views", "Q2 watch time (hours)", "Q1 views (optional)",
                       "Share of views", "Change vs Q1", "What this source tells you"], height=34)
    first, last = 5, 4 + len(TRAFFIC_SOURCES)
    for i, (name, meaning) in enumerate(TRAFFIC_SOURCES):
        r = first + i
        put(ws, f"A{r}", name, bold=True, border=True)
        for col, fm in (("B", INT), ("C", "#,##0.0"), ("D", INT)):
            put(ws, f"{col}{r}", None, color=INPUT_FONT, bg=INPUT_FILL, border=True, fmt=fm)
        put(ws, f"E{r}", f'=IF(OR(B{r}="",SUM($B${first}:$B${last})=0),"",B{r}/SUM($B${first}:$B${last})*100)', border=True, fmt=PCT)
        put(ws, f"F{r}", f'=IF(OR(N(D{r})=0,B{r}=""),"",(B{r}-D{r})/D{r}*100)', border=True, fmt='+0.0"%";-0.0"%";0"%"')
        put(ws, f"G{r}", meaning, color=GREY_TXT, border=True, align=WRAP)
        ws.row_dimensions[r].height = 28
    ws.conditional_formatting.add(f"E{first}:E{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="5B9BD5"))
    change_rules(ws, f"F{first}:F{last}", f"F{first}")
    tr = last + 1  # 17
    put(ws, f"A{tr}", "TOTAL", bold=True, color="FFFFFF", bg=NAVY, border=True)
    for col in "BCD":
        put(ws, f"{col}{tr}", f"=SUM({col}{first}:{col}{last})", bold=True, bg=LIGHT, border=True, fmt=INT)
    put(ws, f"A{tr + 1}", "Top traffic source", bold=True)
    put(ws, f"B{tr + 1}", f'=IF(SUM(B{first}:B{last})=0,"",INDEX(A{first}:A{last},MATCH(MAX(B{first}:B{last}),B{first}:B{last},0)))', bold=True, color=ACCENT)
    put(ws, f"A{tr + 2}", "Its share of views", bold=True)
    put(ws, f"B{tr + 2}", f'=IF(SUM(B{first}:B{last})=0,"",MAX(B{first}:B{last})/SUM(B{first}:B{last})*100)', fmt=PCT, bold=True)

    def block(start, title, labels, note):
        section(ws, f"A{start}:G{start}", title)
        header_row(ws, start + 1, ["Group", "Q2 views / viewers", "Share"], height=28)
        put(ws, f"G{start + 1}", "Note", bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
        s, e = start + 2, start + 1 + len(labels)
        for i, lab in enumerate(labels):
            r = s + i
            put(ws, f"A{r}", lab, bold=True, border=True)
            put(ws, f"B{r}", None, color=INPUT_FONT, bg=INPUT_FILL, border=True, fmt=INT)
            put(ws, f"C{r}", f'=IF(OR(B{r}="",SUM($B${s}:$B${e})=0),"",B{r}/SUM($B${s}:$B${e})*100)', border=True, fmt=PCT)
        ws.conditional_formatting.add(f"C{s}:C{e}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="70AD47"))
        merge_put(ws, f"G{s}:G{e}", note, color=GREY_TXT, align=WRAP)
        return s, e

    vs, ve = block(tr + 4, "Viewer loyalty (Audience tab → New / Casual / Regular viewers)", VIEWER_TYPES,
                   "Casual + Regular = returning viewers. YouTube's 2026 satisfaction-weighted ranking rewards channels people come back to. "
                   "Grow Regulars with series, consistent upload days and community posts.")
    put(ws, f"A{ve + 1}", "Returning viewers share", bold=True)
    put(ws, f"B{ve + 1}", f'=IF(SUM(B{vs}:B{ve})=0,"",(N(B{vs + 1})+N(B{vs + 2}))/SUM(B{vs}:B{ve})*100)', bold=True, fmt=PCT)
    ret_ref = f"B{ve + 1}"
    ds, de = block(ve + 3, "Device type (Advanced mode → Device type)", DEVICES,
                   "High TV share → longer, lean-back videos and readable on-screen text. High mobile share → big, simple thumbnails and text.")
    as_, ae = block(de + 2, "Viewer age (Advanced mode → Viewer age)", AGES,
                    "Use your biggest age group to set tone, references and examples. 25–44 typically skews toward professional, 'use AI at work' content.")
    section(ws, f"A{ae + 2}:G{ae + 2}", "Top countries (Advanced mode → Geography)")
    header_row(ws, ae + 3, ["Country", "Q2 views", "Share"], height=22)
    put(ws, f"G{ae + 3}", "Note", bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
    cs, ce = ae + 4, ae + 8
    for r in range(cs, ce + 1):
        put(ws, f"A{r}", None, color=INPUT_FONT, bg=INPUT_FILL, border=True)
        put(ws, f"B{r}", None, color=INPUT_FONT, bg=INPUT_FILL, border=True, fmt=INT)
        put(ws, f"C{r}", f'=IF(OR(B{r}="",N($B${tr})=0),"",B{r}/$B${tr}*100)',
            border=True, fmt=PCT)
    merge_put(ws, f"G{cs}:G{ce}", "Share is of total Q2 views. If a non-English-speaking country is in your top 5, turn on auto-dubbing (27 languages since Feb 2026) — see Insights.",
              color=GREY_TXT, align=WRAP)
    return ws, {"search_share": "E7", "top_source": f"B{tr + 1}", "top_share": f"B{tr + 2}", "ret": ret_ref,
                "first": first, "last": last}


def build_benchmarks(wb):
    ws = wb.create_sheet("Benchmarks")
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:H1", "Benchmarks — the thresholds behind every grade")
    put(ws, "A2", "Blue numbers are editable assumptions. Change them if your niche differs — every grade in the workbook updates.",
        italic=True, color=GREY_TXT)
    widths(ws, {"A": 34, "B": 14, "C": 12, "D": 12, "E": 12, "F": 50, "G": 60, "H": 16})
    header_row(ws, 4, ["Metric", "Unit", "Average from", "Strong from", "Elite from", "Why it matters", "Source / basis", "Confidence"], height=30)
    for i, (label, unit, a, s, e, why, src, conf) in enumerate(BENCH):
        r = 5 + i
        put(ws, f"A{r}", label, bold=True, border=True, align=WRAP)
        put(ws, f"B{r}", unit, border=True, align=WRAP)
        for col, v in zip("CDE", (a, s, e)):
            put(ws, f"{col}{r}", v, color=INPUT_FONT, bg=INPUT_FILL, border=True, align=CENTER, fmt="0.0")
        put(ws, f"F{r}", why, border=True, align=WRAP)
        put(ws, f"G{r}", src, border=True, align=WRAP, color=GREY_TXT)
        put(ws, f"H{r}", conf, border=True, align=CENTER)
        ws.row_dimensions[r].height = 64
    put(ws, "C4", "Average from", bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
    put(ws, "A13", "Below 'Average from' = Needs work (red)  ·  Average (amber)  ·  Strong (green)  ·  Elite (dark green)", italic=True, color=GREY_TXT)
    put(ws, "A15", "Minimum impressions for a fair CTR comparison", bold=True)
    put(ws, "C15", 1000, color=INPUT_FONT, bg=INPUT_FILL, border=True, fmt=INT)
    put(ws, "D15", "Videos with fewer impressions are skipped when picking your 'best CTR' video (tiny samples give misleading CTRs).",
        italic=True, color=GREY_TXT)
    put(ws, "A17", "Important: benchmarks are averages across very different channels. Your own channel's trend (are you improving month to month?) matters more than any single industry number.",
        italic=True, color=GREY_TXT)
    return ws


def build_dashboard(wb, tr):
    ws = wb.create_sheet("Dashboard", 1)
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 2, "B": 7, "C": 12, "D": 12, "E": 12, "F": 12, "G": 10, "H": 12, "I": 10, "J": 11, "K": 11, "L": 10, "M": 46, "N": 6})
    banner(ws, "B1:M1", "Q2 2026 YouTube Performance Dashboard  ·  Apr 1 – Jun 30")
    merge_put(ws, "B2:M2",
              f'=IF(COUNTA({R("A")})=0,"⚠  NO DATA YET — open the Start Here tab to load your YouTube Studio export. Everything on this page fills in automatically.",'
              f'"✔  "&COUNTA({R("A")})&" videos analysed  ·  "&COUNTIF({R("S")},"Yes")&" published in Q2  ·  "&COUNTIF({R("Q")},"Long")&" long-form, "&COUNTIF({R("Q")},"Short")&" Shorts")',
              bold=True, color=NAVY, bg=LIGHT, align=Alignment(vertical="center", indent=1))
    ws.row_dimensions[2].height = 22

    section(ws, "B4:M4", "Q2 at a glance")
    M = "'Monthly Trend'"

    def q1(col):
        return (f'=IF({M}!{col}10="","Add Q1 in Monthly Trend to compare",'
                f'IF({M}!{col}10>=0,"▲ ","▼ ")&TEXT(ABS({M}!{col}10),"0.0")&"% vs Q1 2026")')

    # ---- layout rows
    TK = 14                       # takeaways section header
    HC = TK + 11                  # health check
    hc_rows = {k: HC + 2 + i for i, k in enumerate(["ctr", "apv_long", "apv_short", "eng", "cpk", "spk", "ret"])}
    LV = HC + 11                  # long vs shorts / mix section header
    lv = {k: LV + 2 + i for i, k in enumerate(["n", "views", "avg", "hours", "apv", "eng", "cpk", "net", "spk", "rpm"])}
    mix_labels = [
        ("WINNER", "WINNER*", "Make sequels, a series, and Shorts from these."),
        ("Great video, weak packaging", "Great video*", "Fastest win: re-test title + thumbnail now."),
        ("Click but leave", "Click but leave*", "Fix the first 30s; make the title match the payoff."),
        ("Topic miss", "Topic miss*", "Retire the angle; don't remake it."),
        ("Average (long-form)", "Average*", "Sharpen hooks and titles."),
        ("Winner Short", "Winner Short*", "Expand into a long-form video."),
        ("Hook problem (Short)", "Hook problem*", "Start on the payoff; no intro."),
        ("Watched but no reaction (Short)", "Watched but no reaction*", "End on a question or hot take."),
    ]
    mix = {lab: LV + 2 + i for i, (lab, _, _) in enumerate(mix_labels)}
    TOP = LV + 13
    TOPS = TOP + 13
    BOT = TOPS + 8
    PAT = BOT + 8
    CH = PAT + 11

    def gsub(tile, key, typical):
        return f'=IF({tile}="","—","Grade: "&$H${hc_rows[key]}&"  ({typical})")'

    tiles = [
        ("B", 5, "VIEWS", f'=IF(N({M}!B8)>0,{M}!B8,SUM({R("E")}))', INT, q1("B")),
        ("E", 5, "WATCH TIME (HOURS)", f'=IF(N({M}!C8)>0,{M}!C8,SUM({R("J")}))', "#,##0", q1("C")),
        ("H", 5, "NET SUBSCRIBERS", f'=IF(N({M}!D8)+N({M}!E8)>0,{M}!F8,SUM({R("X")}))', '+#,##0;-#,##0;0', q1("F")),
        ("K", 5, "EST. REVENUE (USD)", f'=IF(N({M}!K8)>0,{M}!K8,SUM({R("P")}))', '"$"#,##0', q1("K")),
        ("B", 9, "IMPRESSIONS", f'=IF(N({M}!G8)>0,{M}!G8,SUM({R("F")}))', INT, q1("G")),
        ("E", 9, "IMPRESSIONS CTR (LONG-FORM)",
         f'=IFERROR(SUMPRODUCT(({R("Q")}="Long")*{R("F")}*{R("G")})/SUMIFS({R("F")},{R("Q")},"Long",{R("G")},">0"),IF(N({M}!H8)>0,{M}!H8,""))',
         PCT, gsub("E10", "ctr", "typical 4–6%")),
        ("H", 9, "AVG % VIEWED (LONG-FORM)",
         f'=IFERROR(SUMPRODUCT(({R("Q")}="Long")*{R("E")}*{R("I")})/SUMIFS({R("E")},{R("Q")},"Long",{R("I")},">0"),"")', PCT,
         gsub("H10", "apv_long", "good ≈ 30%+")),
        ("K", 9, "ENGAGEMENT RATE", f'=IFERROR((SUM({R("K")})+SUM({R("L")}))/SUM({R("E")})*100,"")', '0.00"%"',
         gsub("K10", "eng", "median ≈ 3.5%")),
    ]
    tile_cells = {}
    for col, r, label, formula, fm, sub in tiles:
        ci = column_index_from_string(col)
        c2 = get_column_letter(ci + 2)
        merge_put(ws, f"{col}{r}:{c2}{r}", label, size=9, bold=True, color=GREY_TXT, bg=PALE, align=CENTER)
        merge_put(ws, f"{col}{r + 1}:{c2}{r + 1}", formula, size=20, bold=True, color=NAVY, bg=PALE, align=CENTER, fmt=fm)
        merge_put(ws, f"{col}{r + 2}:{c2}{r + 2}", sub, size=9, bold=r == 9, color=GREY_TXT, bg=PALE, align=CENTER)
        tile_cells[label] = f"{col}{r + 1}"
        ws.row_dimensions[r + 1].height = 34
        for rr in range(r, r + 3):
            for cc in range(ci, ci + 3):
                ws.cell(row=rr, column=cc).border = Border(left=Side(style="thin", color="FFFFFF"), right=Side(style="thin", color="FFFFFF"))
    grade_rules(ws, "E11:M11", "E11")
    views, hours, net, rev = (tile_cells[k] for k in ("VIEWS", "WATCH TIME (HOURS)", "NET SUBSCRIBERS", "EST. REVENUE (USD)"))
    impr, ctr, apv, eng = (tile_cells[k] for k in ("IMPRESSIONS", "IMPRESSIONS CTR (LONG-FORM)", "AVG % VIEWED (LONG-FORM)", "ENGAGEMENT RATE"))

    # ---- takeaways
    section(ws, f"B{TK}:M{TK}", "Key takeaways — written automatically from your numbers")
    mom = f"{M}!B14"
    best_day, best_len = f"D{PAT + 9}", f"K{PAT + 9}"
    take = [
        f'=IF({ctr}="","• PACKAGING — add impressions + CTR on the Video Data tab to unlock this insight.",'
        f'"• PACKAGING — Long-form CTR was "&TEXT({ctr},"0.0")&"% ("&H{hc_rows["ctr"]}&"; typical channels sit around 4–6%). "&'
        f'IF(OR(H{hc_rows["ctr"]}="Needs work",H{hc_rows["ctr"]}="Average"),"Biggest quick win: run Test & Compare (3 titles + 3 thumbnails) on every new upload and re-package the "&L{mix["Great video, weak packaging"]}&" video(s) flagged Great video, weak packaging.",'
        f'"Packaging is a strength — keep testing bolder variants, and never let a title over-promise."))',

        f'=IF({apv}="","• RETENTION — add Average percentage viewed on the Video Data tab to unlock this insight.",'
        f'"• RETENTION — Long-form viewers watched "&TEXT({apv},"0.0")&"% of each video on average ("&H{hc_rows["apv_long"]}&"; ~30% is good for long-form). "&'
        f'IF(OR(H{hc_rows["apv_long"]}="Needs work",H{hc_rows["apv_long"]}="Average"),"Deliver the promised payoff inside 30 seconds, cut set-up, and add a pattern-break every 60–90s. Study the retention graphs of the "&L{mix["Click but leave"]}&" Click-but-leave video(s) first.",'
        f'"Your content holds attention — lean into the formats behind your "&L{mix["WINNER"]}&" WINNER video(s)."))',

        f'=IFERROR(IF(AND(E{lv["n"]}=0,F{lv["n"]}=0),"• FORMAT — no videos loaded yet.",'
        f'IF(F{lv["n"]}=0,"• FORMAT — no Shorts in Q2. Test 3 Shorts a week cut from your best long-form moments, each linked to the full video.",'
        f'IF(E{lv["n"]}=0,"• FORMAT — Shorts only in Q2. Long-form earns far more per view and builds deeper loyalty; add one long-form video a week.",'
        f'"• FORMAT — Long-form averaged "&TEXT(E{lv["avg"]},"#,##0")&" views/video vs "&TEXT(F{lv["avg"]},"#,##0")&" for Shorts; net subs per 1k views: long-form "&TEXT(E{lv["spk"]},"0.0")&" vs Shorts "&TEXT(F{lv["spk"]},"0.0")&". "&'
        f'IF(F{lv["spk"]}>E{lv["spk"]},"Shorts are your subscriber engine — keep them, but point every Short to a long-form video.","Long-form converts viewers into subscribers better — make it the core, use Shorts as trailers.")))),'
        f'"• FORMAT — add views and subscriber data for both formats to compare.")',

        f'=IFERROR(IF(\'Traffic & Audience\'!{tr["top_source"]}="","• DISCOVERY — fill in the Traffic & Audience tab to see where viewers find you.",'
        f'"• DISCOVERY — "&\'Traffic & Audience\'!{tr["top_source"]}&" drove "&TEXT(\'Traffic & Audience\'!{tr["top_share"]},"0")&"% of views; YouTube search drove "&TEXT(N(\'Traffic & Audience\'!{tr["search_share"]}),"0")&"%. "&'
        f'IF(N(\'Traffic & Audience\'!{tr["search_share"]})<15,"Search is under-used: title one video a week around an exact query people type (e.g. \'how to … with [AI tool]\').",'
        f'"Search is a real asset — group tutorials into playlists so search viewers keep watching.")),"• DISCOVERY — check the Traffic & Audience tab.")',

        f'=IF(AND({best_day}="",{best_len}=""),"• TIMING & LENGTH — add publish dates and durations to unlock this insight.",'
        f'"• TIMING & LENGTH — "&IF({best_day}="","",{best_day}&" uploads averaged the most long-form views. ")&IF({best_len}="","","Best long-form length bucket: "&{best_len}&". ")&'
        f'"Treat these as hypotheses (small samples) and test them on purpose next quarter.")',

        f'=IF({mom}="","• MOMENTUM — add monthly totals on the Monthly Trend tab to see your April→June trajectory.",'
        f'"• MOMENTUM — Views went "&IF({mom}>=0,"up ","down ")&TEXT(ABS({mom}),"0")&"% from April to June. "&'
        f'IF({mom}>=0,"Keep the cadence and topics that drove this — consistency compounds.","Check what changed in May/June (topics, cadence, length); the Diagnosis column on Video Data shows which uploads dragged."))',

        f'=IF(COUNT({R("AC")})=0,"• SCORECARD — load video data to see which videos to double down on.",'
        f'"• SCORECARD — "&L{mix["WINNER"]}&" WINNER video(s) (strong clicks AND strong retention): build sequels and series from them. "&L{mix["Topic miss"]}&" topic miss(es): retire those angles. Every video is diagnosed on the Video Data tab.")',

        "• Q3/Q4 CONTEXT — From Aug 24 2026 YouTube counts a view from the first frame, so post-Q2 view counts will look inflated vs this report. "
        "Judge growth on watch hours, CTR, % viewed and subscribers, not raw views. (Source: 9to5Google, Forbes, Aug 2026.)",
    ]
    for i, t in enumerate(take):
        r = TK + 1 + i
        merge_put(ws, f"B{r}:M{r}", t, align=Alignment(wrap_text=True, vertical="center", indent=1), border=True,
                  bg="FFFFFF" if i % 2 else PALE)
        ws.row_dimensions[r].height = 36

    # ---- health check
    section(ws, f"B{HC}:M{HC}", "Health check vs industry benchmarks (thresholds on the Benchmarks tab)")
    r0 = HC + 1
    merge_put(ws, f"B{r0}:C{r0}", "Metric", bold=True, color="FFFFFF", bg=NAVY, align=CENTER)
    for col, lab in zip("DEFGH", ["Your Q2", "Average from", "Strong from", "Elite from", "Grade"]):
        put(ws, f"{col}{r0}", lab, bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
    merge_put(ws, f"I{r0}:M{r0}", "What to do next", bold=True, color="FFFFFF", bg=NAVY, align=CENTER)
    ws.row_dimensions[r0].height = 28
    short_apv = f'=IFERROR(SUMPRODUCT(({R("Q")}="Short")*{R("E")}*{R("I")})/SUMIFS({R("E")},{R("Q")},"Short",{R("I")},">0"),"")'
    hc_spec = {
        "ctr": ("Impressions CTR (long-form)", f"={ctr}", PCT, [
            "Packaging is the bottleneck. Use Test & Compare on every video; one clear idea per thumbnail (face + 3 words max); make the title a specific outcome or curiosity gap.",
            "In the normal range. Test bolder contrast & emotion in thumbnails and sharper, more specific titles. Re-package older videos that still get impressions.",
            "Above typical. Keep the packaging formula; document what your winning thumbnails share and reuse it.",
            "Elite packaging. Make sure retention keeps up so the click isn't wasted."]),
        "apv_long": ("Avg % viewed (long-form)", f"={apv}", PCT, [
            "Viewers leave early. Open with the payoff/result, cut intros & channel plugs, show-don't-tell demos, pattern-break every 60–90s.",
            "Normal. Find the biggest drop in each retention graph and fix that moment type in future scripts.",
            "Strong. Test slightly longer, deeper videos — satisfied viewers watch more.",
            "Elite. Your structure works — templatise it and scale output."]),
        "apv_short": ("Avg % viewed (Shorts)", short_apv, PCT, [
            "Shorts lose viewers fast. First frame = the most interesting moment; no intro; under 35s until retention improves.",
            "Normal. Tighten pacing and end on a loop that sends viewers back to the start.",
            "Strong. Post more Shorts in this style and link them to related long-form.",
            "Elite. Turn your best Shorts into full long-form deep-dives."]),
        "eng": ("Engagement rate", f"={eng}", '0.00"%"', [
            "Low reaction. Ask one specific question on screen, take a clear stance, pin a comment, and reply to the first hour of comments.",
            "Normal. Add a mid-video prompt ('which tool would you pick?') and reply to comments to seed discussion.",
            "Strong. Mine comments for next video ideas and feature viewers in videos.",
            "Elite. Your community is an asset — consider memberships, a Discord or a newsletter."]),
        "cpk": ("Comments per 1,000 views", f'=IFERROR(SUM({R("L")})/SUM({R("E")})*1000,"")', "0.0", [
            "Very few comments. End with a direct question, and use polls/community posts to warm up the audience.",
            "Normal. Run comment-driven segments ('you asked, I tested').",
            "Strong. Turn top comments into Shorts and follow-ups.",
            "Elite. Make community a pillar (live Q&As, viewer challenges)."]),
        "spk": ("Net subs per 1,000 views", f'=IFERROR({net}/{views}*1000,"")', "0.0", [
            "Views aren't converting. Give a reason to subscribe (a series promise), use end screens to a playlist, and make channel-specific content, not one-offs.",
            "Normal. Add a mid-video 'this is part of a series' cue and a strong channel trailer.",
            "Strong. Keep series formats; promote your best 'binge' playlist.",
            "Elite. Your audience fit is excellent — scale frequency."]),
        "ret": ("Returning viewers share", f"='Traffic & Audience'!{tr['ret']}", PCT, [
            "Mostly one-time viewers. Build recurring series, a consistent upload day and community posts between uploads.",
            "Normal. Create a weekly signature format viewers can expect.",
            "Strong loyalty. Use Hype, premieres and memberships to activate it.",
            "Elite loyalty. Monetise depth: memberships, products, sponsorships."]),
    }
    for key, (label, formula, fm, adv) in hc_spec.items():
        r = hc_rows[key]
        b = BROW[key]
        merge_put(ws, f"B{r}:C{r}", label, bold=True, border=True, align=Alignment(vertical="center", wrap_text=True))
        put(ws, f"D{r}", formula, bold=True, border=True, fmt=fm, align=CENTER)
        for col, bc in zip("EFG", "CDE"):
            put(ws, f"{col}{r}", f"=Benchmarks!{bc}{b}", border=True, fmt="0.0", align=CENTER, color=GREY_TXT)
        put(ws, f"H{r}", f'=IF(D{r}="","—",IF(D{r}>=G{r},"Elite",IF(D{r}>=F{r},"Strong",IF(D{r}>=E{r},"Average","Needs work"))))',
            bold=True, border=True, align=CENTER)
        adv_f = (f'=IF(H{r}="—","Add data to unlock advice.",IF(H{r}="Needs work","{adv[0]}",IF(H{r}="Average","{adv[1]}",'
                 f'IF(H{r}="Strong","{adv[2]}","{adv[3]}"))))')
        merge_put(ws, f"I{r}:M{r}", adv_f, border=True, align=Alignment(wrap_text=True, vertical="center"))
        ws.row_dimensions[r].height = 40
    grade_rules(ws, f"H{HC + 2}:H{HC + 8}", f"H{HC + 2}")

    # ---- long vs shorts
    section(ws, f"B{LV}:G{LV}", "Long-form vs Shorts")
    section(ws, f"I{LV}:M{LV}", "Video scorecard mix (from Diagnosis column)")
    r0 = LV + 1
    merge_put(ws, f"B{r0}:D{r0}", "Metric", bold=True, color="FFFFFF", bg=NAVY, align=CENTER)
    for col, lab in zip("EFG", ["Long-form", "Shorts", "Edge"]):
        put(ws, f"{col}{r0}", lab, bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
    merge_put(ws, f"I{r0}:K{r0}", "Diagnosis", bold=True, color="FFFFFF", bg=NAVY, align=CENTER)
    put(ws, f"L{r0}", "Videos", bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
    put(ws, f"M{r0}", "Next step", bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)

    def by(fmt_name):
        q = f'"{fmt_name}"'
        return {
            "n": f'=COUNTIF({R("Q")},{q})',
            "views": f'=SUMIFS({R("E")},{R("Q")},{q})',
            "avg": "=IFERROR({c}{v}/{c}{n},\"\")",
            "hours": f'=SUMIFS({R("J")},{R("Q")},{q})',
            "apv": f'=IFERROR(SUMPRODUCT(({R("Q")}={q})*{R("E")}*{R("I")})/SUMIFS({R("E")},{R("Q")},{q},{R("I")},">0"),"")',
            "eng": f'=IFERROR((SUMIFS({R("K")},{R("Q")},{q})+SUMIFS({R("L")},{R("Q")},{q}))/SUMIFS({R("E")},{R("Q")},{q})*100,"")',
            "cpk": f'=IFERROR(SUMIFS({R("L")},{R("Q")},{q})/SUMIFS({R("E")},{R("Q")},{q})*1000,"")',
            "net": f'=SUMIFS({R("X")},{R("Q")},{q})',
            "spk": "=IFERROR({c}{net}/{c}{v}*1000,\"\")",
            "rpm": f'=IFERROR(SUMIFS({R("P")},{R("Q")},{q})/SUMIFS({R("E")},{R("Q")},{q})*1000,"")',
        }
    lv_labels = {"n": ("Videos", "0"), "views": ("Total views", INT), "avg": ("Avg views per video", INT),
                 "hours": ("Watch time (hours)", "#,##0"), "apv": ("Avg % viewed (view-weighted)", PCT),
                 "eng": ("Engagement rate", '0.00"%"'), "cpk": ("Comments per 1k views", "0.0"),
                 "net": ("Net subscribers", INT), "spk": ("Net subs per 1k views", "0.0"),
                 "rpm": ("Revenue per 1k views (USD)", MONEY)}
    for col, name in (("E", "Long"), ("F", "Short")):
        fx = by(name)
        for key, r in lv.items():
            v = fx[key].format(c=col, v=lv["views"], n=lv["n"], net=lv["net"])
            put(ws, f"{col}{r}", v, border=True, fmt=lv_labels[key][1], align=CENTER)
    for key, r in lv.items():
        merge_put(ws, f"B{r}:D{r}", lv_labels[key][0], bold=True, border=True)
        if key in ("n", "apv", "views", "hours", "net"):
            put(ws, f"G{r}", "—" if key in ("n", "apv") else
                f'=IF(OR(N(E{r})=0,N(F{r})=0),"",IF(E{r}>F{r},"Long-form",IF(F{r}>E{r},"Shorts","Tie")))',
                border=True, align=CENTER, color=GREY_TXT)
        else:
            put(ws, f"G{r}", f'=IF(OR(E{r}="",F{r}=""),"",IF(E{r}>F{r},"Long-form",IF(F{r}>E{r},"Shorts","Tie")))',
                border=True, align=CENTER, bold=True, color=ACCENT)
    put(ws, f"B{LV + 12}", "Avg % viewed isn't comparable across formats (Shorts loop). Shorts typically earn far less per view than long-form.",
        italic=True, color=GREY_TXT, size=9)

    for (lab, pat, nxt) in mix_labels:
        r = mix[lab]
        merge_put(ws, f"I{r}:K{r}", lab, bold=True, border=True)
        put(ws, f"L{r}", f'=COUNTIF({R("AD")},"{pat}")', border=True, align=CENTER, bold=True)
        put(ws, f"M{r}", nxt, border=True, color=GREY_TXT, align=WRAP)
    for lab, bg, tx in (("WINNER", ELT_BG, ELT_TX), ("Winner Short", ELT_BG, ELT_TX), ("Great video, weak packaging", AMB_BG, AMB_TX),
                        ("Click but leave", AMB_BG, AMB_TX), ("Topic miss", RED_BG, RED_TX), ("Hook problem (Short)", RED_BG, RED_TX),
                        ("Watched but no reaction (Short)", AMB_BG, AMB_TX)):
        ws[f"I{mix[lab]}"].fill = fill(bg)
        ws[f"I{mix[lab]}"].font = f(10, True, tx)

    # ---- top / bottom tables
    def video_table(start, title, n, rank_expr):
        section(ws, f"B{start}:M{start}", title)
        r0 = start + 1
        put(ws, f"B{r0}", "#", bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
        merge_put(ws, f"C{r0}:F{r0}", "Video title", bold=True, color="FFFFFF", bg=NAVY, align=CENTER)
        for col, lab in zip("GHIJKL", ["Format", "Views", "CTR", "% viewed", "Engage-ment", "Net subs"]):
            put(ws, f"{col}{r0}", lab, bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
        put(ws, f"M{r0}", "Diagnosis", bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
        ws.row_dimensions[r0].height = 28
        for k in range(1, n + 1):
            r = r0 + k
            put(ws, f"B{r}", k, border=True, align=CENTER, bold=True)
            put(ws, f"N{r}", f'=IFERROR(MATCH({rank_expr.format(k=k)},{R("AE")},0),"")', color="FFFFFF")

            def pick(col):
                return f'=IF($N{r}="","",IF(INDEX({R(col)},$N{r})="","",INDEX({R(col)},$N{r})))'
            merge_put(ws, f"C{r}:F{r}", pick("A"), border=True, align=Alignment(vertical="center", wrap_text=True))
            for col, src, fm in (("G", "Q", None), ("H", "E", INT), ("I", "G", PCT), ("J", "I", PCT), ("K", "V", '0.00"%"'), ("L", "X", '+#,##0;-#,##0;0')):
                put(ws, f"{col}{r}", pick(src), border=True, fmt=fm, align=CENTER)
            put(ws, f"M{r}", pick("AD"), border=True, align=Alignment(vertical="center", wrap_text=True), size=9)
            ws.row_dimensions[r].height = 30
        rng = f"M{r0 + 1}:M{r0 + n}"
        for word, bg, tx in [("WINNER", ELT_BG, ELT_TX), ("Winner Short", ELT_BG, ELT_TX), ("weak packaging", AMB_BG, AMB_TX),
                             ("Click but leave", AMB_BG, AMB_TX), ("Topic miss", RED_BG, RED_TX), ("Hook problem", RED_BG, RED_TX)]:
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{word}",M{r0 + 1}))'],
                                                           fill=fill(bg), font=Font(name=FONT, color=tx, bold=True, size=9)))
        ws.conditional_formatting.add(f"H{r0 + 1}:H{r0 + n}", DataBarRule(start_type="num", start_value=0, end_type="max", color="5B9BD5"))
        return r0

    top_r0 = video_table(TOP, "Top 10 long-form videos by views — study and repeat these", 10, '"Long-{k}"')
    video_table(TOPS, "Top 5 Shorts by views", 5, '"Short-{k}"')
    nlong = f'COUNTIF({R("AE")},"Long-*")'
    video_table(BOT, "Bottom 5 long-form videos by views — learn from these (shown when you have 10+ long-form videos)", 5,
                f'IF({nlong}<10,"none","Long-"&({nlong}-{{k}}+1))')
    ws.column_dimensions["N"].hidden = True

    # ---- publishing patterns
    section(ws, f"B{PAT}:F{PAT}", "Best publish day (long-form)")
    section(ws, f"H{PAT}:M{PAT}", "Best long-form length")
    r0 = PAT + 1
    for col, lab in zip("BCDEF", ["Day", "Long videos", "Avg views (long)", "Avg CTR (long)", "Shorts / avg views"]):
        put(ws, f"{col}{r0}", lab, bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
    ws.row_dimensions[r0].height = 30
    days = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    for i, d in enumerate(days):
        r = r0 + 1 + i
        put(ws, f"B{r}", d, bold=True, border=True, align=CENTER)
        put(ws, f"C{r}", f'=COUNTIFS({R("T")},B{r},{R("Q")},"Long")', border=True, align=CENTER)
        put(ws, f"D{r}", f'=IFERROR(AVERAGEIFS({R("E")},{R("T")},B{r},{R("Q")},"Long"),"")', border=True, fmt=INT, align=CENTER)
        put(ws, f"E{r}", f'=IFERROR(AVERAGEIFS({R("G")},{R("T")},B{r},{R("Q")},"Long"),"")', border=True, fmt=PCT, align=CENTER)
        put(ws, f"F{r}", f'=IF(COUNTIFS({R("T")},B{r},{R("Q")},"Short")=0,"—",COUNTIFS({R("T")},B{r},{R("Q")},"Short")&" / "&TEXT(AVERAGEIFS({R("E")},{R("T")},B{r},{R("Q")},"Short"),"#,##0"))',
            border=True, align=CENTER)
    d1, d7 = r0 + 1, r0 + 7
    ws.conditional_formatting.add(f"D{d1}:D{d7}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="63BE7B"))
    put(ws, f"B{PAT + 9}", "Best day:", bold=True)
    put(ws, f"D{PAT + 9}", f'=IF(COUNT(D{d1}:D{d7})=0,"",INDEX(B{d1}:B{d7},MATCH(MAX(D{d1}:D{d7}),D{d1}:D{d7},0)))', bold=True, color=ACCENT)

    for col, lab in zip("HIJKLM", ["Length", "", "Videos", "Avg views", "Avg % viewed", "Avg CTR"]):
        put(ws, f"{col}{r0}", lab, bold=True, color="FFFFFF", bg=NAVY, align=CENTER, border=True)
    ws.merge_cells(f"H{r0}:I{r0}")
    buckets = ["Under 8 min", "8–15 min", "15–25 min", "25+ min"]
    for i, bkt in enumerate(buckets):
        r = r0 + 1 + i
        merge_put(ws, f"H{r}:I{r}", bkt, bold=True, border=True, align=CENTER)
        put(ws, f"J{r}", f'=COUNTIF({R("U")},H{r})', border=True, align=CENTER)
        put(ws, f"K{r}", f'=IFERROR(AVERAGEIFS({R("E")},{R("U")},H{r}),"")', border=True, fmt=INT, align=CENTER)
        put(ws, f"L{r}", f'=IFERROR(AVERAGEIFS({R("I")},{R("U")},H{r}),"")', border=True, fmt=PCT, align=CENTER)
        put(ws, f"M{r}", f'=IFERROR(AVERAGEIFS({R("G")},{R("U")},H{r}),"")', border=True, fmt=PCT, align=CENTER)
    b1, b4 = r0 + 1, r0 + 4
    ws.conditional_formatting.add(f"K{b1}:K{b4}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="63BE7B"))
    put(ws, f"H{PAT + 9}", "Best length:", bold=True)
    put(ws, f"K{PAT + 9}", f'=IF(COUNT(K{b1}:K{b4})=0,"",INDEX(H{b1}:H{b4},MATCH(MAX(K{b1}:K{b4}),K{b1}:K{b4},0)))', bold=True, color=ACCENT)

    # ---- charts
    section(ws, f"B{CH}:M{CH}", "Charts")
    mt = wb["Monthly Trend"]
    c1 = BarChart()
    c1.title, c1.y_axis.title = "Views by month", "Views"
    c1.add_data(Reference(mt, min_col=2, min_row=4, max_row=7), titles_from_data=True)
    c1.set_categories(Reference(mt, min_col=1, min_row=5, max_row=7))
    c1.height, c1.width, c1.legend = 7, 12, None
    ws.add_chart(c1, f"B{CH + 1}")
    c2 = BarChart()
    c2.title, c2.y_axis.title = "Net subscribers by month", "Subscribers"
    c2.add_data(Reference(mt, min_col=6, min_row=4, max_row=7), titles_from_data=True)
    c2.set_categories(Reference(mt, min_col=1, min_row=5, max_row=7))
    c2.height, c2.width, c2.legend = 7, 12, None
    c2.series[0].graphicalProperties.solidFill = "70AD47"
    ws.add_chart(c2, f"H{CH + 1}")
    tws = wb["Traffic & Audience"]
    c3 = BarChart()
    c3.type, c3.title = "bar", "Q2 views by traffic source"
    c3.add_data(Reference(tws, min_col=2, min_row=4, max_row=tr["last"]), titles_from_data=True)
    c3.set_categories(Reference(tws, min_col=1, min_row=tr["first"], max_row=tr["last"]))
    c3.height, c3.width, c3.legend = 11, 12, None
    c3.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(c3, f"B{CH + 16}")
    c4 = BarChart()
    c4.type, c4.title = "bar", "Top 10 long-form videos — views"
    c4.add_data(Reference(ws, min_col=8, min_row=top_r0, max_row=top_r0 + 10), titles_from_data=True)
    c4.set_categories(Reference(ws, min_col=2, min_row=top_r0 + 1, max_row=top_r0 + 10))
    c4.height, c4.width, c4.legend = 11, 12, None
    c4.x_axis.scaling.orientation = "maxMin"
    c4.y_axis.title = "Views"
    c4.x_axis.title = "Rank (see table above)"
    ws.add_chart(c4, f"H{CH + 16}")
    ws.freeze_panes = "A4"
    return ws


INSIGHTS = [
    ("YouTube algorithm", "Satisfaction now outweighs raw watch time (2026)",
     "YouTube's ranking leans increasingly on viewer satisfaction — post-view surveys, return visits, 'Not interested' clicks — rather than minutes watched alone. Long-form and Shorts are ranked by separate systems. A shorter video that fully satisfies can beat a padded longer one.",
     "Stop padding videos to hit a length. Optimise for 'did the viewer get what the title promised?'. Your 'Click but leave' videos (Video Data → Diagnosis) are the ones hurting satisfaction most.",
     "Cut every intro to under 20s; state the payoff up front; end with a clear next video so viewers stay on your channel.",
     "Medium", "outlierkit.com/resources/youtube-viewer-satisfaction-algorithm-2026 · meikuio.com/youtube-algorithm-2026"),
    ("Measurement", "Views now count from the first frame (since Aug 24 2026)",
     "YouTube changed public view counting: long-form, live and podcast views now count when playback starts, with no minimum watch time (Shorts already worked this way since Mar 2025). Monetisation still uses engaged views / engaged watch hours.",
     "Your Q3 and Q4 view counts are NOT comparable with this Q2 report — they'll look higher for the same real performance.",
     "Track watch hours, CTR, average % viewed and subscribers as your 'true' growth KPIs in the Q3 Scorecard; read raw views with caution.",
     "High", "9to5google.com/2026/08/17/youtube-view-counts-change · forbes.com (Aug 17 2026) · vidiq.com/blog/post/youtube-view-count-update"),
    ("Policy / AI", "'Inauthentic content' rules target mass-produced AI, not AI use",
     "The YPP 'repetitious content' policy was renamed 'inauthentic content' (Jul 2025) and clarified again in Jul 2026: generic/template content, unsatisfying content and AI personas giving advice on sensitive topics are not monetisable. Channels that use AI as a tool remain eligible.",
     "As an AI-space creator you're fine using AI tools on camera — but avoid templated, faceless, AI-voiced uploads that look interchangeable.",
     "Keep a visible human perspective in every video: your face or voice, real testing, opinions, and original screen recordings.",
     "Medium", "vidiq.com/blog/post/youtube-ai-monetization · air.io/en/monetization/youtube-monetization-policy-changes-2026-a-complete-dated-timeline"),
    ("Packaging", "Test & Compare now tests titles and thumbnails — and picks by watch time",
     "Studio's built-in A/B test runs up to 3 titles and 3 thumbnails on long-form videos for up to 14 days and declares the winner by watch-time share, not clicks. Desktop Studio only; needs Advanced features.",
     "The tool rewards packaging that attracts the right viewer, not just the most clicks — exactly what the satisfaction-weighted algorithm wants.",
     "Run a test on every Q3/Q4 long-form upload and on every 'Great video, weak packaging' video from Q2.",
     "High", "support.google.com/youtube/answer/16391400"),
    ("Benchmarks", "CTR: ~4–6% is typical, 9%+ is top tier",
     "YouTube Help cites 2–10% for most channels. 2026 analyses put the organic average at ~4–6%. CTR naturally falls as YouTube shows a video to colder audiences (more impressions), so a low CTR on a high-impression video can be healthy.",
     "Judge CTR together with impressions. A video at 3% CTR with 500k impressions may be your best video.",
     "Focus packaging fixes on videos with BELOW-average CTR AND low impressions — YouTube stopped testing them.",
     "Medium", "subsub.io/blog/youtube-ctr-what-is-a-good-click-through-rate · focus-digital.co/average-youtube-ctr-organic-paid-benchmarks-2026"),
    ("Benchmarks", "Retention: ~30% average viewed is good for long-form; Shorts ~70%",
     "vidIQ (Jun 2026): good long-form average % viewed ≈ 30%; platform average is in the low-to-mid 20s; very few videos pass 50%. Its Shorts average ~74%. Compare each video with similar-length videos on your own channel.",
     "Use the Retention grade on Video Data; then open each red/amber video's retention graph in Studio to find the exact drop-off moment.",
     "Build a 'drop-off log': for each video note the second where the biggest dip happens and why. Fix the pattern, not the video.",
     "Medium", "vidiq.com/blog/post/increase-audience-retention-youtube"),
    ("Benchmarks", "Engagement: median ~3.5% (likes+comments ÷ views)",
     "A 2026 study of 2,251 videos across 309 channels found a median of 3.48% for long-form and 2.71% for Shorts; ~5% is top-quartile. Channel size barely explained engagement — small channels can be top-tier.",
     "Engagement is a lever you control fully regardless of channel size.",
     "Ask one specific question per video, pin a comment, and reply to comments in the first hour after publishing.",
     "Medium", "overseeros.com/blog/youtube-engagement-rate-benchmarks"),
    ("Format strategy", "Shorts: a funnel, not a business model",
     "An 18,000-channel study (Mar 2025–Mar 2026) found the effect of adding Shorts depends on niche — business/education-style channels that mix Shorts tend to have more subscribers, while some long-form-only niches do better. Shorts RPM is a fraction of long-form.",
     "For an AI channel, Shorts work best as trailers that send viewers to tutorials and deep dives.",
     "Cut 2–3 Shorts from every long-form video, each using 'related video' to link the full version. Judge Shorts on subscribers gained, not views.",
     "Medium", "air.io/en/audience-growth/do-youtube-shorts-help-your-long-form-videos-grow-data-from-18000-channels"),
    ("Growth lever", "Auto-dubbing is open to everyone in 27 languages",
     "Since Feb 2026 YouTube auto-dubbing covers 27 languages with more expressive voices; viewers choose original or dubbed audio, and YouTube says dubbing doesn't hurt discovery.",
     "AI tutorials travel well across languages (screen demos are universal). This is free reach.",
     "Turn on auto-dubbing for your evergreen tutorials; check Traffic & Audience → Top countries after 30 days.",
     "Medium", "storyboard18.com (YouTube expands auto-dubbing to 27 languages) · en.ittefaq.com.bd/15221/youtube-rolls-out-auto-dubbing-globally"),
    ("Growth lever", "Hype boosts small channels (500–500K subscribers)",
     "Viewers can 'Hype' up to 3 long-form videos a week from channels with 500–500K subscribers, within 7 days of upload; smaller channels get bonus points and a weekly leaderboard spot.",
     "If you're in the bracket, your most loyal viewers can push new uploads onto the leaderboard.",
     "Mention Hype in the first 48 hours of each upload (community post + pinned comment), focused on your best video each week.",
     "Medium", "vidiq.com/blog/post/youtube-hype · etcentric.org (Hype now live in 39 countries)"),
    ("AI search", "AI summaries are coming to YouTube and Google",
     "YouTube is testing Gemini-powered AI overviews/highlights in search for Premium users; Google AI Overviews are in ~100 countries; Gemini now draws on YouTube. Creators worry summaries will cut clicks on simple 'answer' videos.",
     "Pure 'what is X' explainers are the most exposed. Hands-on demos, opinions, tests and personality are hard to summarise.",
     "Shift from 'explain' to 'show & judge': real tests, comparisons, builds, verdicts. Add chapters + clear spoken keywords so AI surfaces cite your video.",
     "Medium", "tech.yahoo.com (YouTube is testing its own version of AI Overviews) · startupfortune.com (YouTube embeds AI chatbot into search)"),
    ("AI content niche", "AI tool tutorials & workflows remain the strongest AI niche",
     "2026 niche analyses rank AI tools/workflows among the fastest-growing, highest-CPM clusters (professional, search-driven audiences). Generic AI-voice channels are saturated; hyper-specific 'one tool, one workflow, one outcome' videos win.",
     "Your best content pillar is likely practical, outcome-based tutorials for a specific audience (e.g. 'AI for small business', 'AI for students').",
     "Define 3 content pillars (Tutorials / Tests & Comparisons / News explained) and assign every Q3/Q4 video to one.",
     "Medium", "vidiq.com/blog/post/profitable-ai-niches-youtube · outlierkit.com/resources/articles/top-10-most-profitable-youtube-niches-in-2026"),
    ("AI space (Q3–Q4 2026)", "Personal AI agents went mainstream; models got cheaper and faster",
     "Q3 2026 reports describe a shift from proving agents work to making them general-purpose, with several personal-agent launches in September 2026; new fast, cheap model tiers; growth in open-weight and on-device AI; and more attention on agent security and access controls.",
     "Your audience now asks 'what can I actually delegate to an agent, and is it safe?' — a durable, searchable question set.",
     "Video ideas: 'I let an AI agent run my ___ for a week', 'Best cheap model for ___ (tested)', 'Run AI locally on your laptop', 'Agent safety: 5 settings to change'.",
     "Medium", "finance.biggo.com (Tencent Penguin Intelligence Q3 AI trends report) · agiliway.com/latest-tech-news-q3-2026 · dominotech.net/2026/10/14/ai-trends-q3-roundup"),
    ("AI on YouTube", "Gemini Omni remix and likeness detection",
     "Google I/O 2026 (May) launched Gemini Omni video generation, including Shorts Remix of eligible Shorts (watermarked and linked to the original), and expanded likeness detection to all creators 18+.",
     "Your Shorts can be remixed by others (free reach with a link back), and your face/voice can be protected.",
     "Enable likeness detection in Studio. Test 1 Short per week designed to be remixed (a template or challenge).",
     "Medium", "opus.pro/blog/google-io-2026-ai-video-generation-gemini-updates · medium.com (Everything Google announced at I/O 2026)"),
    ("Strategist view", "In AI content, speed + trust beats volume",
     "AI news moves weekly, but YouTube's satisfaction signals reward videos that genuinely help. The winning pattern is a fast first take (24–72h) on big launches, then a slower, hands-on 'I tested it' follow-up.",
     "Plan a two-beat format: quick 'what it means for you' + in-depth test video one week later.",
     "Keep a 'launch radar' list; pre-script templates so a quality first take ships within 72h.",
     "Strategist judgement", "Synthesis of the sources above + standard creator practice"),
]

ACTIONS = [
    ("P1", "Packaging", "Run Test & Compare (3 titles + 3 thumbnails) on every new long-form upload", "Test & Compare picks by watch time; CTR is the cheapest lever to move", "Impressions CTR (long-form)", "+0.5–1.0 pts vs Q2", "Every upload", "Insights #4, #5"),
    ("P1", "Packaging", "Re-package all Q2 videos diagnosed 'Great video, weak packaging' (new title + thumbnail test)", "Good retention + weak CTR = the content already works; only the 'shop window' is failing", "CTR of those videos", "+20% relative CTR", "Weeks 1–2", "Video Data → Diagnosis"),
    ("P1", "Retention", "Rewrite intros: payoff/result in the first 30s, no channel intro, no 'before we start'", "Satisfaction-weighted ranking; weak openings suppress distribution", "Avg % viewed (long-form)", "+3–5 pts vs Q2", "Every upload", "Insights #1, #6"),
    ("P1", "Retention", "Open the retention graph for every red/amber video; log the biggest drop-off moment and its cause", "Fix the pattern behind the drops rather than single videos", "Drop-off log completed", "All Q2 long-form videos", "Weeks 1–3", "Insights #6"),
    ("P1", "Content", "Make a sequel, series or Shorts from every WINNER video", "Proven topic + proven packaging = lowest-risk upload", "Views per video vs channel avg", "≥ 1.5× channel avg", "Weeks 2–8", "Dashboard → scorecard mix"),
    ("P1", "Measurement", "Track watch hours, CTR, % viewed and subs as the 'true' KPIs (views are inflated after Aug 24 2026)", "View counting changed; raw views aren't comparable with Q2", "Q3 Scorecard filled", "Monthly", "Monthly", "Insights #2"),
    ("P2", "Content strategy", "Define 3 pillars (Tutorials / Tests & Comparisons / AI News Explained) and tag every video", "Specific, outcome-based AI content outperforms generic AI content", "% of uploads mapped to a pillar", "100%", "Week 1", "Insights #12"),
    ("P2", "Content strategy", "Publish 1 search-led tutorial a week titled around an exact query ('How to ___ with ___')", "AI tutorial demand is search-driven and evergreen", "YouTube search share of views", "+5 pts vs Q2", "Weekly", "Traffic & Audience"),
    ("P2", "AI topics", "Cover agents and cheap/fast/local models with a hands-on 'I tested it' angle", "Q3 2026's biggest AI shift; hands-on content resists AI summaries", "Avg views of these videos", "≥ channel avg", "Weeks 2–10", "Insights #11, #13"),
    ("P2", "Shorts", "Cut 2–3 Shorts from each long-form video, each linked via 'related video'", "Shorts work best as a funnel for long-form", "Net subs per 1k Shorts views", "+25% vs Q2", "Every upload", "Insights #8"),
    ("P2", "Community", "Ask one specific question per video, pin a comment, reply during the first hour", "Engagement is size-independent and feeds satisfaction", "Engagement rate / comments per 1k", "≥ 3.5% engagement", "Every upload", "Insights #7"),
    ("P2", "Growth", "Turn on auto-dubbing for the top 10 evergreen tutorials", "Free reach in 27 languages; YouTube says it doesn't hurt discovery", "Views from non-primary-language countries", "+10% vs Q2", "Week 2", "Insights #9"),
    ("P3", "Growth", "Ask loyal viewers to Hype new uploads (if 500–500K subs)", "Hype gives small channels a leaderboard boost", "Hype leaderboard appearances", "1+ per month", "Weekly", "Insights #10"),
    ("P3", "Policy / brand", "Audit for 'inauthentic' risk: no templated AI-voice uploads; keep a visible human perspective", "Protects monetisation under the 2025–26 YPP clarifications", "Audit done", "Done", "Week 1", "Insights #3"),
    ("P3", "Protection", "Enable likeness detection; design 1 Short a week to be remixable", "I/O 2026 features: protect your face and voice, and earn remix reach", "Feature enabled; remixes", "Enabled", "Week 1", "Insights #14"),
    ("P3", "Operations", "Keep an AI 'launch radar' and a 72-hour first-take script template", "Speed + trust pattern for AI news", "Time from launch to upload", "≤ 72h", "Ongoing", "Insights #15"),
]


def build_insights(wb):
    ws = wb.create_sheet("Insights & Research")
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:H1", "Insights & Research — YouTube in 2026, the creator economy and the AI space")
    put(ws, "A2", "Research gathered Oct 2026. Confidence: High = official or several consistent reports · Medium = several secondary sources · "
                  "Strategist judgement = my professional synthesis. Verify specific figures before quoting them publicly.",
        italic=True, color=GREY_TXT)
    widths(ws, {"A": 5, "B": 16, "C": 30, "D": 58, "E": 44, "F": 44, "G": 13, "H": 40})
    header_row(ws, 4, ["#", "Theme", "Finding", "What's happening (research)", "What it means for you", "Do this", "Confidence", "Sources"], height=26)
    for i, (theme, title, what, means, do, conf, src) in enumerate(INSIGHTS):
        r = 5 + i
        vals = [i + 1, theme, title, what, means, do, conf, src]
        for ci, v in enumerate(vals):
            col = get_column_letter(ci + 1)
            put(ws, f"{col}{r}", v, border=True, align=WRAP if ci not in (0, 6) else CENTER,
                bold=ci in (0, 2), color=GREY_TXT if ci == 7 else ("000000" if ci != 2 else NAVY),
                bg=PALE if i % 2 == 0 else None)
        ws.row_dimensions[r].height = 120
    end = 4 + len(INSIGHTS)
    for word, bg, tx in [("High", GRN_BG, GRN_TX), ("Medium", AMB_BG, AMB_TX), ("Strategist", LIGHT, NAVY)]:
        ws.conditional_formatting.add(f"G5:G{end}", FormulaRule(formula=[f'ISNUMBER(SEARCH("{word}",G5))'],
                                                                 fill=fill(bg), font=Font(name=FONT, bold=True, color=tx)))
    ws.freeze_panes = "D5"
    return ws


def build_actions(wb):
    ws = wb.create_sheet("Q3 Action Plan")
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:K1", "Q3 Action Plan — prioritised, measurable, research-backed")
    put(ws, "A2", "P1 = do first (highest impact, lowest effort) · P2 = core growth bets · P3 = supporting moves. Q3 has closed (today is Oct 2026): "
                  "mark what you did in Q3, then carry open items into Q4. Update Status and Owner in the yellow cells.",
        italic=True, color=GREY_TXT)
    widths(ws, {"A": 5, "B": 9, "C": 15, "D": 52, "E": 44, "F": 24, "G": 18, "H": 13, "I": 22, "J": 14, "K": 14})
    header_row(ws, 4, ["#", "Priority", "Area", "Action", "Why (evidence)", "KPI", "Target", "When", "Linked insight / data", "Owner", "Status"], height=26)
    for i, (pri, area, act, why, kpi, tgt, when, link) in enumerate(ACTIONS):
        r = 5 + i
        vals = [i + 1, pri, area, act, why, kpi, tgt, when, link]
        for ci, v in enumerate(vals):
            col = get_column_letter(ci + 1)
            put(ws, f"{col}{r}", v, border=True, align=CENTER if ci in (0, 1) else WRAP, bold=ci in (1, 3))
        put(ws, f"J{r}", "Me", color=INPUT_FONT, bg=INPUT_FILL, border=True, align=CENTER)
        put(ws, f"K{r}", "Not started", color=INPUT_FONT, bg=INPUT_FILL, border=True, align=CENTER)
        ws.row_dimensions[r].height = 48
    end = 4 + len(ACTIONS)
    dv = DataValidation(type="list", formula1='"Not started,In progress,Done,Dropped"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K5:K{end}")
    for word, bg, tx in [("P1", RED_BG, RED_TX), ("P2", AMB_BG, AMB_TX), ("P3", LIGHT, NAVY)]:
        ws.conditional_formatting.add(f"B5:B{end}", FormulaRule(formula=[f'B5="{word}"'], fill=fill(bg), font=Font(name=FONT, bold=True, color=tx)))
    for word, bg, tx in [("Done", GRN_BG, GRN_TX), ("In progress", AMB_BG, AMB_TX), ("Not started", "F2F2F2", "595959"), ("Dropped", "D9D9D9", "7F7F7F")]:
        ws.conditional_formatting.add(f"K5:K{end}", FormulaRule(formula=[f'K5="{word}"'], fill=fill(bg), font=Font(name=FONT, bold=True, color=tx)))
    put(ws, f"D{end + 2}", "Progress", bold=True)
    put(ws, f"E{end + 2}", f'=COUNTIF(K5:K{end},"Done")&" of "&COUNTA(D5:D{end})&" actions done ("&TEXT(COUNTIF(K5:K{end},"Done")/COUNTA(D5:D{end}),"0%")&")"',
        bold=True, color=ACCENT)
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:K{end}"
    return ws


def build_scorecard(wb, tr):
    ws = wb.create_sheet("Q3 Scorecard")
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:H1", "Q3 Scorecard — Q2 baseline → Q3 target → Q3 actual")
    put(ws, "A2", "Q2 baseline fills from the Dashboard. Edit the target change (yellow) and enter Q3 actuals (Jul 1 – Sep 30 2026) to grade the quarter.",
        italic=True, color=GREY_TXT)
    widths(ws, {"A": 34, "B": 14, "C": 10, "D": 13, "E": 14, "F": 14, "G": 13, "H": 44})
    header_row(ws, 4, ["Metric", "Q2 2026 actual", "Change type", "Target change", "Q3 target", "Q3 actual", "% of target", "Note"], height=30)
    D = "Dashboard"
    rows = [
        ("Watch time (hours)", f"={D}!E6", "%", 15, "#,##0", "Best 'true' growth KPI after the Aug 2026 view-count change."),
        ("Net subscribers", f"={D}!H6", "%", 20, INT, ""),
        ("Impressions", f"={D}!B10", "%", 15, INT, "Rises when YouTube trusts your videos (satisfaction)."),
        ("Impressions CTR, long-form (%)", f"={D}!E10", "pts", 0.5, PCT, "Target is added in percentage points (e.g. 5.0% → 5.5%)."),
        ("Avg % viewed, long-form (%)", f"={D}!H10", "pts", 3, PCT, ""),
        ("Engagement rate (%)", f"={D}!K10", "pts", 0.5, '0.00"%"', ""),
        ("Views", f"={D}!B6", "%", 15, INT, "Caution: Q3 views count from the first frame after Aug 24 2026 — expect inflation."),
        ("Long-form videos published", "='Monthly Trend'!L8", "%", 0, "0", "Consistency beats bursts — set a number you can sustain."),
        ("Shorts published", "='Monthly Trend'!M8", "%", 0, "0", ""),
        ("Est. revenue (USD)", f"={D}!K6", "%", 15, MONEY, ""),
    ]
    for i, (label, base, kind, chg, fm, note) in enumerate(rows):
        r = 5 + i
        put(ws, f"A{r}", label, bold=True, border=True)
        put(ws, f"B{r}", base, border=True, fmt=fm, align=CENTER)
        put(ws, f"C{r}", kind, border=True, align=CENTER, color=GREY_TXT)
        put(ws, f"D{r}", chg, color=INPUT_FONT, bg=INPUT_FILL, border=True, align=CENTER,
            fmt='+0"%";-0"%";0"%"' if kind == "%" else '+0.0" pts";-0.0" pts"')
        put(ws, f"E{r}", f'=IF(N(B{r})=0,"",IF(C{r}="pts",B{r}+D{r},B{r}*(1+D{r}/100)))', border=True, fmt=fm, align=CENTER, bold=True)
        put(ws, f"F{r}", None, color=INPUT_FONT, bg=INPUT_FILL, border=True, fmt=fm, align=CENTER)
        put(ws, f"G{r}", f'=IF(OR(F{r}="",N(E{r})=0),"",F{r}/E{r}*100)', border=True, fmt='0"%"', align=CENTER, bold=True)
        put(ws, f"H{r}", note, color=GREY_TXT, border=True, align=WRAP)
        ws.row_dimensions[r].height = 30
    end = 4 + len(rows)
    ws.conditional_formatting.add(f"G5:G{end}", FormulaRule(formula=['AND(ISNUMBER(G5),G5>=100)'], fill=fill(GRN_BG), font=Font(name=FONT, bold=True, color=GRN_TX)))
    ws.conditional_formatting.add(f"G5:G{end}", FormulaRule(formula=['AND(ISNUMBER(G5),G5>=90,G5<100)'], fill=fill(AMB_BG), font=Font(name=FONT, bold=True, color=AMB_TX)))
    ws.conditional_formatting.add(f"G5:G{end}", FormulaRule(formula=['AND(ISNUMBER(G5),G5<90)'], fill=fill(RED_BG), font=Font(name=FONT, bold=True, color=RED_TX)))
    put(ws, f"A{end + 2}", "Green ≥ 100% of target · Amber 90–99% · Red < 90%. Default targets are moderate (+15% growth, +0.5 pt CTR); raise them if Q2 was a slow quarter.",
        italic=True, color=GREY_TXT)
    return ws


# ======================================================================= demo data
def fill_demo(wb):
    rnd = random.Random(42)
    ws = wb["Video Data"]
    long_titles = [
        "I Let an AI Agent Run My Inbox for 7 Days", "5 AI Tools That Replaced My Entire Workflow", "Claude vs ChatGPT vs Gemini: Real Work Test",
        "Build a Website With AI in 20 Minutes (No Code)", "The AI News You Actually Need This Week", "How I Use AI to Plan My Whole Week",
        "Run AI Locally on Your Laptop: Full Setup", "I Tried Every AI Video Generator So You Don't Have To", "AI Agents Explained Simply",
        "Automate Your Job With AI (Step-by-Step)", "The Best Free AI Tools in 2026", "I Built an App With AI and Zero Coding Skills",
        "Prompting Masterclass: 10 Techniques That Work", "Is AI Coming for YouTubers?", "AI for Students: Study 2x Faster",
        "My Honest Review of the New AI Glasses", "Make Money With AI: What Actually Works", "AI Spreadsheet Tricks Nobody Shows You",
        "How AI Chooses What You Watch", "Testing AI Voice Clones (Scary Good)", "The AI Tools I Pay For (and Why)", "Q2 AI Recap: What Mattered",
    ]
    short_titles = [
        "This AI trick saves me 2 hours", "AI wrote my email… and it worked", "Stop using ChatGPT like this", "One prompt to rule them all",
        "AI agent ordered my groceries", "Free AI tool you need", "AI vs me: who's faster?", "This AI video isn't real",
        "3 AI shortcuts for Excel", "Local AI in 30 seconds", "The prompt that fixed my code", "AI made this song in 10s",
        "Don't trust AI with this", "Gemini's hidden feature", "Claude's best use case", "AI study hack", "This changed how I work", "Wait for the end 🤯",
    ]
    days = [date(2026, 4, 1) + timedelta(days=i) for i in range(91)]
    vids = []
    for t in long_titles:
        dur = rnd.choice([420, 540, 660, 780, 900, 1080, 1320, 1620])
        impr = int(rnd.lognormvariate(11.3, 0.7))
        ctr = round(max(1.8, rnd.gauss(5.2, 1.6)), 1)
        views = int(impr * ctr / 100 * rnd.uniform(1.05, 1.4))
        apv = round(max(14, min(58, rnd.gauss(31, 8))), 1)
        avd_s = int(dur * apv / 100)
        vids.append(dict(t=t, d=rnd.choice(days), fmt="Long", dur=dur, v=views, i=impr, ctr=ctr, apv=apv, avd=avd_s))
    for t in short_titles:
        dur = rnd.randint(18, 58)
        impr = int(rnd.lognormvariate(10.8, 0.9))
        views = int(impr * rnd.uniform(0.35, 0.75))
        apv = round(max(38, min(112, rnd.gauss(72, 16))), 1)
        ctr = round(rnd.uniform(1.5, 5.0), 1)
        vids.append(dict(t=t, d=rnd.choice(days), fmt="Short", dur=dur, v=views, i=impr, ctr=ctr, apv=apv, avd=int(dur * apv / 100)))
    vids.sort(key=lambda x: x["d"])
    for k, vd in enumerate(vids):
        r = FIRST + k
        long_ = vd["fmt"] == "Long"
        likes = int(vd["v"] * rnd.uniform(0.018, 0.05 if long_ else 0.04))
        comments = int(vd["v"] * rnd.uniform(0.0008, 0.006 if long_ else 0.002))
        shares = int(vd["v"] * rnd.uniform(0.002, 0.008))
        sg = int(vd["v"] * rnd.uniform(0.002, 0.012 if long_ else 0.006))
        sl = int(sg * rnd.uniform(0.05, 0.25))
        wh = round(vd["v"] * vd["avd"] / 3600, 1)
        rev = round(vd["v"] / 1000 * (rnd.uniform(6, 14) if long_ else rnd.uniform(0.05, 0.25)), 2)
        h, rem = divmod(vd["avd"], 3600)
        m, s = divmod(rem, 60)
        row = [vd["t"], vd["d"], vd["fmt"] if k % 3 else None, vd["dur"], vd["v"], vd["i"], vd["ctr"], time(h, m, s),
               vd["apv"], wh, likes, comments, shares, sg, sl, rev]
        for ci, val in enumerate(row):
            ws.cell(row=r, column=ci + 1).value = val
    mt = wb["Monthly Trend"]
    monthly = [(412000, 21800, 2950, 410, 3_900_000, 5.1, time(0, 3, 48), 118000, 2150),
               (468000, 24900, 3320, 455, 4_350_000, 5.3, time(0, 3, 55), 131000, 2480),
               (521000, 27600, 3790, 470, 4_820_000, 5.0, time(0, 3, 51), 142000, 2760)]
    for i, vals in enumerate(monthly):
        r = 5 + i
        for col, v in zip(["B", "C", "D", "E", "G", "H", "I", "J", "K"], vals):
            mt[f"{col}{r}"].value = v
    for col, v in zip(["B", "C", "D", "E", "G", "H", "I", "J", "K"], (1_190_000, 61500, 8100, 1300, 11_800_000, 4.8, time(0, 3, 40), 300000, 5900)):
        mt[f"{col}9"].value = v
    tws = wb["Traffic & Audience"]
    tvals = [(540000, 33100, 470000), (310000, 19800, 280000), (205000, 11200, 150000), (290000, 3900, 230000),
             (18000, 1100, 16000), (21000, 900, 15000), (9000, 600, 9000), (6500, 500, 5000), (3500, 200, 3000),
             (8000, 400, 7000), (2500, 100, 2000), (1500, 50, 1000)]
    for i, (v, wh, q1) in enumerate(tvals):
        r = 5 + i
        tws[f"B{r}"].value, tws[f"C{r}"].value, tws[f"D{r}"].value = v, wh, q1
    for r, v in zip(range(23, 26), (262000, 98000, 31000)):
        tws[f"B{r}"].value = v
    for r, v in zip(range(30, 35), (780000, 420000, 160000, 34000, 1000)):
        tws[f"B{r}"].value = v
    for r, v in zip(range(38, 45), (21000, 290000, 520000, 360000, 150000, 50000, 14000)):
        tws[f"B{r}"].value = v
    for r, (c, v) in zip(range(48, 53), (("United States", 520000), ("India", 210000), ("United Kingdom", 95000), ("Germany", 61000), ("Brazil", 48000))):
        tws[f"A{r}"].value, tws[f"B{r}"].value = c, v
    for name in wb.sheetnames:
        s = wb[name]
        s.sheet_properties.tabColor = "C00000"


def main():
    demo = "--demo" in sys.argv
    wb = Workbook()
    build_start(wb, demo)
    build_video_data(wb)
    build_monthly(wb)
    _, tr = build_traffic(wb)
    build_benchmarks(wb)
    build_dashboard(wb, tr)
    build_insights(wb)
    build_actions(wb)
    build_scorecard(wb, tr)
    order = ["Start Here", "Dashboard", "Video Data", "Monthly Trend", "Traffic & Audience",
             "Insights & Research", "Q3 Action Plan", "Q3 Scorecard", "Benchmarks"]
    wb._sheets = [wb[n] for n in order]
    if not demo:
        colors = {"Start Here": NAVY, "Dashboard": ACCENT, "Video Data": "FFC000", "Monthly Trend": "FFC000",
                  "Traffic & Audience": "FFC000", "Insights & Research": "70AD47", "Q3 Action Plan": "70AD47",
                  "Q3 Scorecard": "70AD47", "Benchmarks": "A6A6A6"}
        for n, c in colors.items():
            wb[n].sheet_properties.tabColor = c
    else:
        fill_demo(wb)
        for n in order:
            if n not in ("Start Here",):
                wb[n]["A1" if n not in ("Dashboard",) else "B1"].value = "[SAMPLE DATA] " + wb[n]["A1" if n != "Dashboard" else "B1"].value
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    wb.active = 1 if demo else 0
    wb.active.sheet_view.tabSelected = True
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 90
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = "YouTube_Q2_2026_Review_SAMPLE.xlsx" if demo else "YouTube_Q2_2026_Review.xlsx"
    wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
