"""Load a YouTube Studio Advanced-mode export into the review workbook.

    python import_studio_export.py YouTube_Q2_2026_Review.xlsx "Table data.csv"

Columns are matched by YouTube's own header names, so the order in the CSV
doesn't matter and missing metrics are simply left blank. Only the yellow input
columns (A–P) of the 'Video Data' tab are written; every formula is kept.
Open the workbook in Excel afterwards — it recalculates on open.
"""
import csv
import sys
from datetime import datetime, time

from openpyxl import load_workbook

FIRST, LAST = 5, 204

# workbook column -> accepted CSV headers (lower-case)
COLUMNS = {
    "A": ["video title", "title"],
    "B": ["video publish time", "publish time", "published"],
    "C": ["format", "content type"],
    "D": ["duration"],
    "E": ["views"],
    "F": ["impressions"],
    "G": ["impressions click-through rate (%)", "impressions ctr (%)", "impressions click-through rate"],
    "H": ["average view duration"],
    "I": ["average percentage viewed (%)", "average percentage viewed"],
    "J": ["watch time (hours)"],
    "K": ["likes"],
    "L": ["comments added", "comments"],
    "M": ["shares"],
    "N": ["subscribers gained"],
    "O": ["subscribers lost"],
    "P": ["estimated revenue (usd)", "estimated revenue"],
}


def number(v):
    v = (v or "").strip().replace(",", "").replace("%", "").replace("$", "")
    if v in ("", "-", "—"):
        return None
    try:
        n = float(v)
    except ValueError:
        return None
    return int(n) if n.is_integer() else n


def duration_value(v):
    v = (v or "").strip()
    if not v:
        return None
    if ":" in v:
        parts = [int(float(p)) for p in v.split(":")]
        while len(parts) < 3:
            parts.insert(0, 0)
        h, m, s = parts
        return time(min(h, 23), m, s)
    secs = number(v)
    if secs is None:
        return None
    secs = int(secs)
    return time(min(secs // 3600, 23), (secs % 3600) // 60, secs % 60)


def parse_date(v):
    v = (v or "").strip()
    for candidate in (v, v[:10]):
        for fmt in ("%b %d, %Y", "%B %d, %Y", "%Y-%m-%d", "%d %b %Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(candidate, fmt).date()
            except ValueError:
                continue
    return None


def fmt_value(v):
    v = (v or "").strip().lower()
    if v.startswith("short"):
        return "Short"
    if v in ("video", "videos", "long", "long-form", "live", "live stream"):
        return "Long"
    return None


def main(xlsx, csv_path):
    with open(csv_path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit("The CSV has no rows.")
    headers = {h.strip().lower(): h for h in rows[0].keys()}
    mapping = {}
    for col, names in COLUMNS.items():
        for n in names:
            if n in headers:
                mapping[col] = headers[n]
                break
    if "A" not in mapping:
        sys.exit("Couldn't find a 'Video title' column. Export from Advanced mode → Content tab.")

    # Studio puts a "Total" row first (it has no video title) — skip it.
    data = [r for r in rows if (r.get(mapping["A"]) or "").strip() and (r.get(mapping["A"]) or "").strip().lower() != "total"]
    if len(data) > LAST - FIRST + 1:
        print(f"Note: {len(data)} videos found; only the first {LAST - FIRST + 1} (by CSV order) fit the sheet.")
        data = data[:LAST - FIRST + 1]

    wb = load_workbook(xlsx)
    ws = wb["Video Data"]
    for r in range(FIRST, LAST + 1):
        for col in COLUMNS:
            ws[f"{col}{r}"].value = None
    for i, row in enumerate(data):
        r = FIRST + i
        for col, header in mapping.items():
            raw = row.get(header, "")
            if col == "A":
                val = raw.strip()
            elif col == "B":
                val = parse_date(raw)
            elif col == "C":
                val = fmt_value(raw)
            elif col == "H":
                val = duration_value(raw)
            else:
                val = number(raw)
            ws[f"{col}{r}"].value = val
    wb.save(xlsx)
    found = ", ".join(sorted(mapping[c] for c in mapping))
    missing = [COLUMNS[c][0] for c in COLUMNS if c not in mapping and c != "C"]
    print(f"Loaded {len(data)} videos into '{xlsx}'.\nMatched columns: {found}")
    if missing:
        print("Not in this export (left blank): " + ", ".join(missing))
    print("Open the workbook in Excel — the Dashboard recalculates automatically.")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
